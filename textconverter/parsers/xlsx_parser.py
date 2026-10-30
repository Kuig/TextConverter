"""Excel (.xlsx) parser: reads a workbook with openpyxl into the AST.

Each visible sheet is a level 1 heading followed by its table (a workbook with
one sheet has no heading). The first row of the sheet is the table's header and
the data starts at A1, so the cell coordinates in the formulas stay true. A cell
holds its value (a formula cell its stored result), a comment as a footnote, a
hyperlink as a link and an image anchored to it as an image. The formulas come
as an ``excel`` code block under the table (see ``_excel_formulas``). Merged
cells, charts, pivot tables, filters and conditional formatting are dropped,
with a single warning listing them.
"""
from __future__ import annotations

import datetime
from collections import Counter

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from ..ast import CodeBlock, Document, Footnote, Heading, Image, Link, Node, Table, TableCell, TableRow, Text
from ..images import ImageStore
from ..logger import log_warning
from ._excel_formulas import FORMULA_LANGUAGE, compress


def format_value(value: object) -> str:
    """A cell value as text: integers without ``.0``, dates in ISO form, booleans as TRUE/FALSE."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else format(value, ".15g")
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, datetime.time):
        return value.isoformat(timespec="seconds")
    return str(value)


class _Style:
    """The formatting of a run of text."""

    def __init__(self, bold=False, italic=False, strike=False, script=None):
        self.bold, self.italic, self.strike, self.script = bold, italic, strike, script

    @classmethod
    def of(cls, font, base: _Style | None = None) -> _Style:
        """The style of an openpyxl font (cell or inline), on top of ``base``."""
        base = base or cls()
        align = getattr(font, "vertAlign", None)
        return cls(bool(font.b) or base.bold, bool(font.i) or base.italic, bool(font.strike) or base.strike,
                   {"superscript": "sup", "subscript": "sub"}.get(align) or base.script)

    def text(self, content: str) -> Text:
        return Text(content, bold=self.bold, italic=self.italic, strikethrough=self.strike, script=self.script)


class _Reader:
    def __init__(self, workbook, values, output_dir: str | None, image_dir_name: str | None, write_images: bool):
        self.workbook = workbook
        self.values = values
        self.write_images = write_images
        self.dropped: Counter[str] = Counter()
        self.uncached = 0
        self._store = ImageStore(output_dir, image_dir_name)

    def read(self) -> Document:
        sheets = []
        for ws in self.workbook.worksheets:
            if ws.sheet_state != "visible":
                continue
            read = self._sheet(ws, self.values[ws.title])
            if read:
                sheets.append((ws.title, *read))
        self.dropped["chart sheets"] += len(self.workbook.chartsheets)
        children: list[Node] = []
        for title, table, formulas in sheets:
            if len(sheets) > 1:
                children.append(Heading(level=1, children=[Text(title)]))
            children.append(table)
            if formulas:
                children.append(CodeBlock("\n".join(formulas), language=FORMULA_LANGUAGE))
        doc = Document(children=children)
        props = self.workbook.properties
        if props.title:
            doc.title = [Text(props.title)]
        if props.creator:
            doc.author = [Text(props.creator)]
        if props.created:
            doc.date = [Text(props.created.date().isoformat())]
        what = ", ".join(f"{name} ({count})" if count > 1 else name for name, count in self.dropped.items() if count)
        if what:
            log_warning(f"Excel features without a counterpart were dropped: {what}.")
        if self.uncached:
            log_warning(f"{self.uncached} formula cell(s) have no stored result and were read as empty; "
                        "open and save the workbook in Excel to store them.")
        return doc

    # ---- sheets -----------------------------------------------------------

    def _sheet(self, ws, ws_values) -> tuple[Table, list[str]] | None:
        self.dropped["merged cells"] += len(ws.merged_cells.ranges)
        self.dropped["charts"] += len(ws._charts)
        self.dropped["pivot tables"] += len(ws._pivots)
        self.dropped["filters"] += 1 if ws.auto_filter.ref else 0
        self.dropped["conditional formatting"] += len(ws.conditional_formatting)

        cells: dict[tuple[int, int], list[Node]] = {}
        formulas: dict[tuple[int, int], str] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.__class__.__name__ == "MergedCell":
                    continue
                nodes = self._cell(cell, ws_values.cell(cell.row, cell.column).value, formulas)
                if nodes:
                    cells[(cell.row, cell.column)] = nodes
        if self.write_images:
            self._images(ws, cells)
        if not cells:
            return None

        height = max(r for r, _ in cells)
        width = max(c for _, c in cells)
        grid = [[TableCell(cells.get((r, c), [])) for c in range(1, width + 1)] for r in range(1, height + 1)]
        # Cell formatting is not read on the header row: the writer makes it bold.
        for cell in grid[0]:
            cell.children = [_plain(n) for n in cell.children]
        return Table(headers=grid[0], rows=[TableRow(r) for r in grid[1:]]), compress(formulas)

    def _cell(self, cell, cached, formulas: dict[tuple[int, int], str]) -> list[Node]:
        value = cell.value
        if isinstance(value, (ArrayFormula, DataTableFormula)):
            self.dropped["array formulas"] += 1
            value = cached
        elif cell.data_type == "f" and isinstance(value, str):
            formulas[(cell.row, cell.column)] = value
            self.uncached += cached is None
            value = cached
        base = _Style.of(cell.font)
        if isinstance(value, CellRichText):
            nodes = [_run(part, base) for part in value]
        else:
            nodes = [base.text(format_value(value))] if value is not None and format_value(value) else []
        nodes = [n for n in nodes if n is not None]
        link = cell.hyperlink
        target = link and (link.target or (f"#{link.location}" if link.location else None))
        if nodes and target:
            nodes = [Link(url=target, content=nodes)]
        if cell.comment and cell.comment.text.strip():
            nodes.append(Footnote([Text(cell.comment.text.strip())]))
        return nodes

    def _images(self, ws, cells: dict[tuple[int, int], list[Node]]) -> None:
        for image in ws._images:
            anchor = getattr(image.anchor, "_from", None)
            if anchor is None:
                self.dropped["floating images"] += 1
                continue
            try:
                src = self._store.save(image._data(), image.format or "png")
            except Exception:
                self.dropped["unreadable images"] += 1
                continue
            cells.setdefault((anchor.row + 1, anchor.col + 1), []).append(Image(src=src))


def _run(part: str | TextBlock, base: _Style) -> Text | None:
    """One run of a rich text cell, with its own formatting on top of the cell's."""
    if isinstance(part, TextBlock):
        content, style = part.text, _Style.of(part.font if isinstance(part.font, InlineFont) else InlineFont(), base)
    else:
        content, style = part, base
    return style.text(content) if content else None


def _plain(node: Node) -> Node:
    """The node without text formatting (inside a link too)."""
    if isinstance(node, Text):
        return Text(node.content)
    if isinstance(node, Link):
        node.content = [_plain(n) for n in node.content]
    return node


def parse_xlsx(
    file_path: str,
    *,
    output_dir: str | None = None,
    image_dir_name: str | None = None,
    write_images: bool = True,
) -> Document:
    """Parse an Excel workbook into an AST Document.

    Args:
        file_path: Path of the ``.xlsx`` file.
        output_dir: Base directory for the extracted images (current directory if None).
        image_dir_name: Image subdirectory name inside ``output_dir``.
        write_images: Extract the images anchored to cells; when False they are ignored.

    Returns:
        A Document with a table per visible sheet (a heading each when there are several).
    """
    workbook = load_workbook(file_path, rich_text=True)
    values = load_workbook(file_path, data_only=True)
    try:
        return _Reader(workbook, values, output_dir, image_dir_name, write_images).read()
    finally:
        workbook.close()
        values.close()
