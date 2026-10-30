"""Word (.docx) parser: reads a document with python-docx into the AST.

Styles are recognized by name (python-docx maps Word's built-in names to their
English form), never by style id, which Word localizes. Features Word has and
the AST does not (underline, highlight, headers and footers, shapes) are
dropped, with a single warning listing them.
"""
from __future__ import annotations

import hashlib
import os
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field, replace

from lxml import etree

from ..ast import (
    Abstract, BlockQuote, CodeBlock, CodeInline, Document, Equation, Footnote, Heading, HorizontalRule, Image, Label,
    LineBreak, Link, ListBlock, ListItem, Node, Paragraph, Table, TableCell, TableRow, Text,
    abstract_from,
)
from ..images import store_image
from ..logger import log_warning
from ..renderers._common import text_runs
from ._latex_macros import lift_macros
from ._math import split_math_env
from ._omml import omml_to_latex

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {
    "w": W, "m": M, "r": R,
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "v": "urn:schemas-microsoft-com:vml",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}

_OFFICE = "urn:schemas-microsoft-com:office:office"

# Fonts whose text is code; the lower-cased family name.
_MONOSPACE = frozenset({
    "consolas", "courier new", "courier", "lucida console", "cascadia code", "cascadia mono", "menlo",
    "monaco", "source code pro", "dejavu sans mono", "liberation mono", "andale mono", "fira code",
    "fira mono", "jetbrains mono", "roboto mono", "ubuntu mono", "sf mono", "noto sans mono",
    "lucida sans typewriter", "monospace",
})
_CODE_STYLES = frozenset({"code", "source code", "preformatted text", "html preformatted", "macro text"})
_QUOTE_STYLES = frozenset({"quote", "intense quote", "block text"})
_HEADING = re.compile(r"^Heading (\d)$")
_LIST_STYLE = re.compile(r"^List (Bullet|Number)(?: (\d))?$")
_TASK = {"☐": False, "☑": True}
_HYPERLINK_INSTR = re.compile(r'HYPERLINK\s+(?:"([^"]*)")?(?:.*?\\l\s+"([^"]*)")?', re.IGNORECASE)
_LABEL_SEPARATOR = re.compile(r"^\s*[:.\-\u2013\u2014]?\s*")
# Notes that Word keeps in the notes part but are layout, not content.
_NOTE_LAYOUT = ("separator", "continuationSeparator", "continuationNotice")


def _w(name: str) -> str:
    return f"{{{W}}}{name}"


def _xpath(el: etree._Element, expression: str) -> list:
    """Evaluate an XPath expression with the prefixes of ``NS`` (python-docx's own ``xpath`` has fixed ones)."""
    return etree.XPath(expression, namespaces=NS)(el)


def _local(el: etree._Element) -> str:
    return el.tag.rsplit("}", 1)[-1]


def _flag(el: etree._Element | None) -> bool:
    """An on/off run property: present, and not switched off with ``w:val``."""
    return el is not None and el.get(_w("val"), "1").lower() not in ("0", "false", "off")


def _val(el: etree._Element | None, default: str | None = None) -> str | None:
    return el.get(_w("val"), default) if el is not None else default


@dataclass
class _RunFormat:
    bold: bool = False
    italic: bool = False
    strike: bool = False
    script: str | None = None
    code: bool = False
    equation: bool = False
    underline: bool = False
    highlight: bool = False


@dataclass
class _Field:
    """A field (or a hyperlink) being read: the nodes of its result collect in ``out``."""
    instr: str = ""
    separated: bool = False
    link: str | None = None
    out: list[Node] = field(default_factory=list)


@dataclass
class _Inline:
    """The state of reading one paragraph's inline content."""
    out: list[Node] = field(default_factory=list)
    stack: list[_Field] = field(default_factory=list)
    extra: list[Node] = field(default_factory=list)
    seq_end: int | None = None

    def emit(self, node: Node) -> None:
        (self.stack[-1].out if self.stack else self.out).append(node)


@dataclass
class _ListRef:
    """Where a list paragraph sits: the list it belongs to (``key``), its level and kind."""
    key: object
    level: int
    ordered: bool
    start: int = 1


class _Numbering:
    """The list definitions of ``numbering.xml``: kind and first number of each list level."""

    def __init__(self, root: etree._Element | None):
        self._levels: dict[str, dict[int, tuple[str, int]]] = {}
        self._nums: dict[str, tuple[str, dict[int, tuple[str | None, int | None]]]] = {}
        if root is None:
            return
        for abstract in root.findall(_w("abstractNum")):
            self._levels[abstract.get(_w("abstractNumId"))] = {
                int(lvl.get(_w("ilvl"), 0)): (_val(lvl.find(_w("numFmt")), "decimal"),
                                              int(_val(lvl.find(_w("start")), "1")))
                for lvl in abstract.findall(_w("lvl"))
            }
        for num in root.findall(_w("num")):
            overrides = {}
            for over in num.findall(_w("lvlOverride")):
                lvl = over.find(_w("lvl"))
                fmt = _val(lvl.find(_w("numFmt"))) if lvl is not None else None
                start = _val(over.find(_w("startOverride")))
                if start is None and lvl is not None:
                    start = _val(lvl.find(_w("start")))
                overrides[int(over.get(_w("ilvl"), 0))] = (fmt, int(start) if start else None)
            self._nums[num.get(_w("numId"))] = (_val(num.find(_w("abstractNumId")), ""), overrides)

    def level(self, num_id: str, level: int) -> tuple[bool, int]:
        """``(ordered, first number)`` of a list level; an unknown list is a bullet list."""
        abstract_id, overrides = self._nums.get(num_id, ("", {}))
        fmt, start = self._levels.get(abstract_id, {}).get(level, ("bullet", 1))
        over_fmt, over_start = overrides.get(level, (None, None))
        fmt = over_fmt or fmt
        return fmt not in ("bullet", "none"), over_start if over_start is not None else start


def parse_docx(
    file_path: str,
    *,
    output_dir: str | None = None,
    image_dir_name: str | None = None,
    write_images: bool = True,
) -> Document:
    """Parse a Word document into an AST Document.

    Args:
        file_path: Path of the ``.docx`` file.
        output_dir: Base directory for the extracted images (current directory if None).
        image_dir_name: Image subdirectory name inside ``output_dir``.
        write_images: Extract the embedded images; when False they are left out.

    Raises:
        ValueError: If the file is not a readable ``.docx`` document.
    """
    import docx
    from docx.opc.exceptions import PackageNotFoundError
    try:
        document = docx.Document(file_path)
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, etree.XMLSyntaxError) as exc:
        raise ValueError(f"Not a readable Word (.docx) document: {file_path} ({exc})") from exc
    return _Reader(document, output_dir, image_dir_name, write_images).read()


class _Reader:
    """Reads one python-docx document; see ``parse_docx``."""

    def __init__(self, document, output_dir: str | None, image_dir_name: str | None, write_images: bool):
        from docx.enum.style import WD_STYLE_TYPE
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        self.document = document
        self.part = document.part
        self._style_types = WD_STYLE_TYPE
        self._rt = RT
        self.image_dir_name = image_dir_name
        self.target_dir = os.path.join(output_dir or os.getcwd(), image_dir_name or "")
        self.write_images = write_images
        self.dropped: Counter[str] = Counter()
        self._images: dict[str, str] = {}
        self._styles: dict[tuple[str, str | None], object] = {}
        self._notes: dict[str, tuple[object, dict[str, etree._Element]]] = {}
        self._numbering: _Numbering | None = None
        self._title: list[Node] = []

    # ---- document -------------------------------------------------------

    def read(self) -> Document:
        body = self.document.element.body
        children = self._blocks(body, top=True)
        props = self.document.core_properties
        doc = Document(children=children, title=self._title)
        if not doc.title and props.title:
            doc.title = [Text(content=props.title)]
        if props.author:
            doc.author = [Text(content=props.author)]
        if props.created:
            doc.date = [Text(content=props.created.date().isoformat())]
        self._note_headers_and_footers()
        if self.dropped:
            what = ", ".join(f"{name} ({count})" if count > 1 else name for name, count in self.dropped.items())
            log_warning(f"Word features without a counterpart were dropped: {what}.")
        return doc

    def _note_headers_and_footers(self) -> None:
        for kind in ("header", "footer"):
            for section in self.document.sections:
                part = getattr(section, kind)
                if not part.is_linked_to_previous and any(p.text.strip() for p in part.paragraphs):
                    self.dropped[kind] += 1
                    break

    # ---- styles ---------------------------------------------------------

    def _style(self, style_id: str | None, kind: str):
        key = (kind, style_id)
        if key not in self._styles:
            self._styles[key] = self.document.styles.get_by_id(style_id, getattr(self._style_types, kind))
        return self._styles[key]

    def _paragraph_style(self, p: etree._Element):
        pstyle = p.find(f"{_w('pPr')}/{_w('pStyle')}")
        return self._style(_val(pstyle), "PARAGRAPH")

    def _style_numbering(self, style) -> tuple[str, int] | None:
        """The list a paragraph style applies (``numId``, level), looking through its base styles."""
        while style is not None:
            pr = style.element.pPr
            if pr is not None and pr.numPr is not None and pr.numPr.numId is not None:
                level = pr.numPr.ilvl.val if pr.numPr.ilvl is not None else 0
                return str(pr.numPr.numId.val), int(level)
            style = style.base_style
        return None

    def _list_ref(self, p: etree._Element, style) -> _ListRef | None:
        """The list the paragraph belongs to, or None."""
        num_pr = p.find(f"{_w('pPr')}/{_w('numPr')}")
        if num_pr is not None:
            num_id = _val(num_pr.find(_w("numId")), "0")
            if num_id == "0":
                return None
            level = int(_val(num_pr.find(_w("ilvl")), "0"))
            ordered, start = self._numbering_model().level(num_id, level)
            return _ListRef(num_id, level, ordered, start)
        named = _LIST_STYLE.match(style.name or "") if style is not None else None
        if named:
            return _ListRef(("style", named.group(1)), int(named.group(2) or 1) - 1, named.group(1) == "Number")
        numbered = self._style_numbering(style)
        if numbered and numbered[0] != "0":
            ordered, start = self._numbering_model().level(*numbered)
            return _ListRef(numbered[0], numbered[1], ordered, start)
        return None

    def _numbering_model(self) -> _Numbering:
        if self._numbering is None:
            try:
                root = etree.fromstring(self.document.part.part_related_by(self._rt.NUMBERING).blob)
            except KeyError:
                root = None
            self._numbering = _Numbering(root)
        return self._numbering

    # ---- blocks ---------------------------------------------------------

    @staticmethod
    def _body_elements(parent: etree._Element):
        """The paragraphs and tables of a body or text box, through content controls."""
        for child in parent:
            tag = child.tag if isinstance(child.tag, str) else ""
            if tag in (_w("p"), _w("tbl")):
                yield child
            elif tag == _w("sdt"):
                content = child.find(_w("sdtContent"))
                if content is not None:
                    yield from _Reader._body_elements(content)

    def _blocks(self, parent: etree._Element, top: bool = False) -> list[Node]:
        """The blocks of a body or text box."""
        flow = _Flow(self)
        elements = list(self._body_elements(parent))
        for index, el in enumerate(elements):
            if el.tag == _w("tbl"):
                flow.table(self._table(el))
                continue
            style = self._paragraph_style(el)
            name = (style.name if style is not None else "") or ""
            lowered = name.lower()
            nodes, ctx = self._inline_paragraph(el)
            followed_by_table = index + 1 < len(elements) and elements[index + 1].tag == _w("tbl")
            list_ref = self._list_ref(el, style)
            if lowered == "caption" and nodes:
                flow.caption(nodes, ctx, followed_by_table)
            elif list_ref is not None:
                flow.list_item(list_ref, self._flatten(nodes, ctx))
            elif lowered == "title" and top and not self._title and nodes:
                flow.close()
                self._title = nodes
                flow.out.extend(ctx.extra)
            elif lowered == "equation display" and (code := self._plain_text(el).strip()):
                body, env = split_math_env(code)
                flow.leaf(Equation(code=body, inline=False, env=env))
            elif _HEADING.match(name) and nodes:
                flow.leaf(Heading(level=min(int(_HEADING.match(name).group(1)), 6), children=nodes), ctx.extra)
            elif lowered in _QUOTE_STYLES and nodes:
                flow.group("quote", Paragraph(children=nodes), ctx.extra)
            elif lowered == "abstract" and nodes:
                flow.group("abstract", Paragraph(children=nodes), ctx.extra)
            elif lowered in _CODE_STYLES or self._is_code(nodes):
                flow.group("code", self._plain_text(el, keep_tabs=True).rstrip(), ctx.extra)
            elif nodes:
                flow.leaf(Paragraph(children=nodes), ctx.extra)
            else:
                flow.close()
                if el.find(f"{_w('pPr')}/{_w('pBdr')}/{_w('bottom')}") is not None:
                    flow.out.append(HorizontalRule())
                flow.out.extend(ctx.extra)
        flow.close()
        return flow.out

    def _flatten(self, nodes: list[Node], ctx: _Inline) -> list[Node]:
        """The inline nodes of a paragraph that cannot hold blocks: display equations stay in the text."""
        extras = []
        for node in ctx.extra:
            if isinstance(node, Equation):
                extras.append(replace(node, inline=True))
            else:
                self.dropped["text box"] += 1
        return [*nodes, *extras]

    @staticmethod
    def _is_code(nodes: list[Node]) -> bool:
        if not any(isinstance(n, CodeInline) and n.code.strip() for n in nodes):
            return False
        return all(isinstance(n, (CodeInline, LineBreak)) or (isinstance(n, Text) and not n.content.strip())
                   for n in nodes)

    def _plain_text(self, p: etree._Element, keep_tabs: bool = False) -> str:
        """The text of a paragraph's runs, ``\\n`` for a line break; tabs stay when ``keep_tabs``."""
        parts = []
        for run in p.iter(_w("r")):
            if any(a.tag == _w("del") for a in run.iterancestors()):
                continue
            for child in run:
                tag = _local(child)
                if tag == "t":
                    parts.append(child.text or "")
                elif tag == "tab":
                    parts.append("\t" if keep_tabs else " ")
                elif tag in ("br", "cr") and child.get(_w("type")) not in ("page", "column"):
                    parts.append("\n")
        return "".join(parts)

    # ---- tables ---------------------------------------------------------

    def _table(self, tbl: etree._Element) -> Table:
        rows = []
        for tr in tbl.findall(_w("tr")):
            cells: list[TableCell] = []
            for tc in tr.findall(_w("tc")):
                props = tc.find(_w("tcPr"))
                merge = props.find(_w("vMerge")) if props is not None else None
                continued = merge is not None and _val(merge, "continue") != "restart"
                span = int(_val(props.find(_w("gridSpan")), "1")) if props is not None else 1
                cells.append(TableCell(children=[] if continued else self._cell(tc)))
                cells.extend(TableCell() for _ in range(span - 1))
            rows.append((tr, cells))
        table = Table()
        if rows and self._has_header(tbl, rows[0][0]):
            table.headers = rows[0][1]
            rows = rows[1:]
            for cell in table.headers:
                # A header is already set apart; its bold text would be bold twice.
                cell.children = [replace(n, bold=False) if isinstance(n, Text) else n for n in cell.children]
        table.rows = [TableRow(cells=cells) for _, cells in rows]
        return table

    @staticmethod
    def _has_header(tbl: etree._Element, first_row: etree._Element) -> bool:
        """Whether the first row is a header: marked as a repeating row, or the table's Header Row option."""
        if first_row.find(f"{_w('trPr')}/{_w('tblHeader')}") is not None:
            return True
        look = tbl.find(f"{_w('tblPr')}/{_w('tblLook')}")
        if look is None:
            return False
        if look.get(_w("firstRow")) is not None:
            return look.get(_w("firstRow")).lower() in ("1", "true", "on")
        try:
            return bool(int(look.get(_w("val"), "0"), 16) & 0x20)
        except ValueError:
            return False

    def _cell(self, tc: etree._Element) -> list[Node]:
        """A cell's content as inline nodes, one line per paragraph (a nested table is flattened)."""
        paragraphs = []
        for child in tc:
            if child.tag == _w("p"):
                paragraphs.append(child)
            elif child.tag == _w("tbl"):
                paragraphs.extend(child.iter(_w("p")))
        out: list[Node] = []
        for p in paragraphs:
            nodes, ctx = self._inline_paragraph(p)
            nodes = self._flatten(nodes, ctx)
            if nodes:
                out.extend([LineBreak()] if out else [])
                out.extend(nodes)
        return out

    # ---- inline content -------------------------------------------------

    def _inline_paragraph(self, p: etree._Element) -> tuple[list[Node], _Inline]:
        """A paragraph's inline nodes, with what it carries besides them (display equations, text boxes)."""
        ctx = _Inline()
        self._inline_children(p, ctx)
        while ctx.stack:
            self._end_field(ctx)
        nodes = lift_macros(text_runs(ctx.out))
        return self._tidy(nodes), ctx

    @staticmethod
    def _tidy(nodes: list[Node]) -> list[Node]:
        """Merge adjacent code and trim the paragraph's edges."""
        merged: list[Node] = []
        for node in nodes:
            if isinstance(node, CodeInline) and merged and isinstance(merged[-1], CodeInline):
                merged[-1] = CodeInline(code=merged[-1].code + node.code)
            else:
                merged.append(node)
        while merged and isinstance(merged[0], LineBreak):
            del merged[0]
        while merged and isinstance(merged[-1], LineBreak):
            del merged[-1]
        if merged and isinstance(merged[0], Text):
            merged[0] = replace(merged[0], content=merged[0].content.lstrip())
        if merged and isinstance(merged[-1], Text):
            merged[-1] = replace(merged[-1], content=merged[-1].content.rstrip())
        return [n for n in merged if not (isinstance(n, Text) and not n.content)]

    def _inline_children(self, parent: etree._Element, ctx: _Inline) -> None:
        previous_equation_run = False
        for child in parent:
            if not isinstance(child.tag, str):
                continue
            tag = child.tag
            was_equation_run = False
            if tag == _w("r"):
                was_equation_run = self._run(child, ctx, previous_equation_run)
            elif tag == _w("hyperlink"):
                self._hyperlink(child, ctx)
            elif tag in (_w("ins"), _w("smartTag"), _w("customXml"), _w("moveTo"), _w("dir"), _w("bdo")):
                self._inline_children(child, ctx)
            elif tag == _w("sdt"):
                content = child.find(_w("sdtContent"))
                if content is not None:
                    self._inline_children(content, ctx)
            elif tag == _w("fldSimple"):
                ctx.stack.append(_Field(instr=child.get(_w("instr"), ""), separated=True))
                self._inline_children(child, ctx)
                self._end_field(ctx)
            elif tag == f"{{{M}}}oMath":
                ctx.emit(Equation(code=omml_to_latex(child), inline=True))
            elif tag == f"{{{M}}}oMathPara":
                ctx.extra.extend(Equation(code=omml_to_latex(m), inline=False)
                                 for m in child.findall(f"{{{M}}}oMath"))
            previous_equation_run = was_equation_run

    def _hyperlink(self, el: etree._Element, ctx: _Inline) -> None:
        url = self._link_target(el.get(f"{{{R}}}id"))
        anchor = el.get(_w("anchor"))
        if anchor:
            url = f"{url or ''}#{anchor}"
        ctx.stack.append(_Field(link=url, separated=True))
        self._inline_children(el, ctx)
        self._end_field(ctx)

    def _link_target(self, rel_id: str | None) -> str | None:
        rel = self.part.rels.get(rel_id) if rel_id else None
        return rel.target_ref if rel is not None and rel.is_external else None

    def _format(self, run: etree._Element) -> _RunFormat:
        fmt = _RunFormat()
        props = run.find(_w("rPr"))
        if props is None:
            return fmt
        fmt.bold = _flag(props.find(_w("b")))
        fmt.italic = _flag(props.find(_w("i")))
        fmt.strike = _flag(props.find(_w("strike"))) or _flag(props.find(_w("dstrike")))
        fmt.script = {"superscript": "sup", "subscript": "sub"}.get(_val(props.find(_w("vertAlign")), ""))
        underline = props.find(_w("u"))
        fmt.underline = (underline is not None and _val(underline, "single") != "none"
                         and run.getparent().tag != _w("hyperlink"))
        fmt.highlight = _val(props.find(_w("highlight")), "none") != "none"
        fonts = props.find(_w("rFonts"))
        if fonts is not None:
            names = (fonts.get(_w(a)) for a in ("ascii", "hAnsi", "cs"))
            fmt.code = any(n and n.lower() in _MONOSPACE for n in names)
        char_style = props.find(_w("rStyle"))
        if char_style is not None:
            style = self._style(_val(char_style), "CHARACTER")
            name = (style.name if style is not None else "") or ""
            fmt.equation = name == "Equation"
            fmt.bold = fmt.bold or name in ("Strong", "Intense Emphasis")
            fmt.italic = fmt.italic or name in ("Emphasis", "Intense Emphasis")
        return fmt

    def _run(self, run: etree._Element, ctx: _Inline, after_equation: bool) -> bool:
        """Read a run into ``ctx``; True when it was an equation run (which the next one extends)."""
        fmt = self._format(run)
        text: list[str] = []

        def flush() -> None:
            if not text:
                return
            content = "".join(text)
            text.clear()
            if fmt.equation:
                last = (ctx.stack[-1].out if ctx.stack else ctx.out)
                if after_equation and last and isinstance(last[-1], Equation):
                    last[-1] = Equation(code=last[-1].code + content, inline=True)
                else:
                    ctx.emit(Equation(code=content, inline=True))
                return
            if content.strip() and fmt.underline:
                self.dropped["underlined text"] += 1
            if content.strip() and fmt.highlight:
                self.dropped["highlighted text"] += 1
            if fmt.code:
                ctx.emit(CodeInline(code=content))
            else:
                ctx.emit(Text(content=content.replace("\t", " "), bold=fmt.bold, italic=fmt.italic,
                              strikethrough=fmt.strike, script=fmt.script))

        self._run_children(run, ctx, text, flush)
        flush()
        return fmt.equation

    def _run_children(self, run: etree._Element, ctx: _Inline, text: list[str], flush) -> None:
        for child in run:
            if not isinstance(child.tag, str):
                continue
            tag = _local(child)
            if tag == "t":
                text.append(child.text or "")
            elif tag == "tab":
                text.append("\t")
            elif tag == "noBreakHyphen":
                text.append("-")
            elif tag in ("br", "cr"):
                if child.get(_w("type")) not in ("page", "column"):
                    flush()
                    ctx.emit(LineBreak())
            elif tag == "fldChar":
                flush()
                self._field_char(child, ctx)
            elif tag == "instrText":
                if ctx.stack:
                    ctx.stack[-1].instr += child.text or ""
            elif tag in ("footnoteReference", "endnoteReference"):
                flush()
                note = self._note(child, "footnote" if tag == "footnoteReference" else "endnote")
                if note is not None:
                    ctx.emit(note)
            elif tag == "drawing":
                flush()
                self._drawing(child, ctx)
            elif tag == "pict":
                flush()
                self._pict(child, ctx)
            elif tag == "AlternateContent":
                flush()
                branch = child.find(f"{{{NS['mc']}}}Choice")
                if branch is None:
                    branch = child.find(f"{{{NS['mc']}}}Fallback")
                if branch is not None:
                    self._run_children(branch, ctx, text, flush)
            elif tag == "object":
                self.dropped["embedded object"] += 1

    # ---- fields ---------------------------------------------------------

    def _field_char(self, el: etree._Element, ctx: _Inline) -> None:
        kind = el.get(_w("fldCharType"))
        if kind == "begin":
            ctx.stack.append(_Field())
        elif kind == "separate" and ctx.stack:
            ctx.stack[-1].separated = True
        elif kind == "end" and ctx.stack:
            self._end_field(ctx)

    def _end_field(self, ctx: _Inline) -> None:
        """Close the innermost field: a hyperlink wraps its result, a SEQ number marks the caption label."""
        fld = ctx.stack.pop()
        instr = fld.instr.strip()
        kind = instr.split(None, 1)[0].upper() if instr else ""
        url = fld.link
        if kind == "HYPERLINK":
            m = _HYPERLINK_INSTR.match(instr)
            url = (m.group(1) or "") + (f"#{m.group(2)}" if m and m.group(2) else "") if m else None
        if url and fld.out:
            ctx.emit(Link(url=url, content=fld.out))
        else:
            for node in fld.out:
                ctx.emit(node)
        if kind == "SEQ" and not ctx.stack:
            ctx.seq_end = len(ctx.out)

    # ---- notes, images, text boxes --------------------------------------

    def _note(self, ref: etree._Element, kind: str) -> Footnote | None:
        if kind not in self._notes:
            rel = self._rt.FOOTNOTES if kind == "footnote" else self._rt.ENDNOTES
            try:
                part = self.document.part.part_related_by(rel)
                root = etree.fromstring(part.blob)
                self._notes[kind] = (part, {
                    n.get(_w("id")): n for n in root.findall(_w(kind)) if n.get(_w("type")) not in _NOTE_LAYOUT
                })
            except KeyError:
                self._notes[kind] = (None, {})
        notes_part, notes = self._notes[kind]
        note = notes.get(ref.get(_w("id")))
        if note is None:
            return None
        main_part, self.part = self.part, notes_part
        try:
            content: list[Node] = []
            for p in note.findall(_w("p")):
                nodes, ctx = self._inline_paragraph(p)
                nodes = self._flatten(nodes, ctx)
                if nodes:
                    content.extend([LineBreak()] if content else [])
                    content.extend(nodes)
        finally:
            self.part = main_part
        return Footnote(content=content)

    def _image(self, rel_id: str | None, alt: str) -> Image | None:
        if not self.write_images:
            return None
        try:
            part = self.part.related_parts[rel_id]
        except KeyError:
            self.dropped["linked or missing image"] += 1
            return None
        digest = hashlib.sha1(part.blob).hexdigest()
        if digest not in self._images:
            ext = re.sub(r"[^a-z0-9]", "", part.partname.ext.lower()) or "bin"
            if ext in ("emf", "wmf"):
                log_warning(f"Image in {ext.upper()} format kept as a file, but most viewers cannot show it.")
            basename = os.path.basename(store_image(part.blob, self.target_dir, f"image_{len(self._images) + 1}.{ext}"))
            self._images[digest] = f"{self.image_dir_name}/{basename}" if self.image_dir_name else basename
        return Image(src=self._images[digest], alt=alt)

    def _drawing(self, drawing: etree._Element, ctx: _Inline) -> None:
        outer = "not(ancestor::w:txbxContent)"
        blips = _xpath(drawing, f".//a:blip[@r:embed][{outer}]")
        boxes = _xpath(drawing, f".//w:txbxContent[{outer}]")
        props = drawing.find(".//wp:docPr", NS)
        alt = (props.get("descr") or props.get("title") or "") if props is not None else ""
        for blip in blips:
            image = self._image(blip.get(f"{{{R}}}embed"), alt)
            if image is not None:
                ctx.emit(image)
        self._text_boxes(boxes, ctx)
        if not blips and not boxes:
            self.dropped["shape"] += 1

    def _pict(self, pict: etree._Element, ctx: _Inline) -> None:
        images = _xpath(pict, ".//v:imagedata[@r:id]")
        boxes = _xpath(pict, ".//w:txbxContent[not(ancestor::w:txbxContent)]")
        for data in images:
            alt = data.getparent().get("alt") or data.get(f"{{{_OFFICE}}}title") or ""
            image = self._image(data.get(f"{{{R}}}id"), alt)
            if image is not None:
                ctx.emit(image)
        self._text_boxes(boxes, ctx)
        if not images and not boxes and any(isinstance(c.tag, str) and _local(c) != "shapetype" for c in pict):
            self.dropped["shape"] += 1

    def _text_boxes(self, boxes: list[etree._Element], ctx: _Inline) -> None:
        for box in boxes:
            blocks = self._blocks(box)
            if blocks:
                ctx.extra.append(BlockQuote(children=blocks))


class _Flow:
    """Collects the blocks of a body: groups consecutive quote, abstract and code paragraphs, builds lists."""

    def __init__(self, reader: _Reader):
        self.reader = reader
        self.out: list[Node] = []
        self._group: tuple[str, list] | None = None
        self._lists: list[tuple[ListBlock, object, int]] = []
        self._seen: Counter = Counter()
        self._caption: tuple[list[Node], str | None] | None = None

    def close(self) -> None:
        """End the open group or list."""
        if self._group is not None:
            kind, items = self._group
            self._group = None
            if kind == "code":
                code = "\n".join(items).strip("\n")
                if code.strip():
                    self.out.append(CodeBlock(code=code))
            elif kind == "quote":
                self.out.append(BlockQuote(children=items))
            else:
                self.out.append(abstract_from(items) or Abstract(children=items))
        self._lists = []

    def group(self, kind: str, item, extra: list[Node]) -> None:
        if self._group is not None and self._group[0] != kind:
            self.close()
        self._lists = []
        if self._group is None:
            self._group = (kind, [])
        self._group[1].append(item)
        if extra:
            self.close()
            self.out.extend(extra)

    def leaf(self, node: Node, extra: list[Node] = ()) -> None:
        self.close()
        self.out.append(node)
        self.out.extend(extra)

    def table(self, table: Table) -> None:
        self.close()
        if self._caption is not None:
            table.caption, table.label = self._caption[0], self._caption[1]
            self._caption = None
        self.out.append(table)

    def caption(self, nodes: list[Node], ctx: _Inline, before_table: bool) -> None:
        """A caption belongs to the table that follows it, else to the table or figure just before it."""
        self.close()
        text = nodes
        if ctx.seq_end:
            text = self.reader._tidy(lift_macros(text_runs(ctx.out[ctx.seq_end:])))
            if text and isinstance(text[0], Text):
                text[0] = replace(text[0], content=_LABEL_SEPARATOR.sub("", text[0].content, count=1))
        label = next((n.name for n in text if isinstance(n, Label)), None)
        text = self.reader._tidy([n for n in text if not isinstance(n, Label)])
        previous = self.out[-1] if self.out else None
        figure = self._lone_image(previous)
        if before_table:
            self._caption = (text, label)
        elif isinstance(previous, Table) and not previous.caption:
            previous.caption, previous.label = text, label
        elif figure is not None and not figure.caption:
            figure.caption, figure.label = text, label
        else:
            self.out.append(Paragraph(children=nodes))
        self.out.extend(ctx.extra)

    @staticmethod
    def _lone_image(node: Node | None) -> Image | None:
        """The image a paragraph holds on its own, else None."""
        if not isinstance(node, Paragraph):
            return None
        content = [n for n in node.children if not (isinstance(n, Text) and not n.content.strip())]
        return content[0] if len(content) == 1 and isinstance(content[0], Image) else None

    def list_item(self, ref: _ListRef, content: list[Node]) -> None:
        if self._group is not None:
            self.close()
        while self._lists and self._lists[-1][2] > ref.level:
            self._lists.pop()
        if self._lists and self._lists[-1][2] == ref.level and self._lists[-1][1] != ref.key:
            self._lists.pop()
        for key in [k for k in self._seen if k[0] == ref.key and k[1] > ref.level]:
            del self._seen[key]
        checked = None
        if content and isinstance(content[0], Text) and content[0].content[:1] in _TASK:
            checked = _TASK[content[0].content[0]]
            content[0] = replace(content[0], content=content[0].content[1:].lstrip())
        if self._lists and self._lists[-1][2] == ref.level:
            block = self._lists[-1][0]
        else:
            seen = self._seen[(ref.key, ref.level)]
            block = ListBlock(ordered=ref.ordered, start=ref.start + seen if ref.ordered else 1)
            if self._lists and self._lists[-1][0].items:
                self._lists[-1][0].items[-1].children.append(block)
            else:
                self.out.append(block)
            self._lists.append((block, ref.key, ref.level))
        block.items.append(ListItem(children=[Paragraph(children=content)], checked=checked))
        self._seen[(ref.key, ref.level)] += 1
