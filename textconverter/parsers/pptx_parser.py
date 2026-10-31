"""PowerPoint (.pptx) parser: reads a presentation with python-pptx into the AST.

A section is a level 1 heading and a slide a level 2 one (a title starting with
``#`` is a deeper slide: ``# Title`` is level 3). A slide's main text box is its
body; every other text box becomes a quote whose first line gives its frame.
Speaker notes become a footnote on the slide's heading. The theme, animations
and transitions are ignored; features the AST cannot hold (videos, shapes
without text, ...) are dropped, with a single warning listing them.
"""
from __future__ import annotations

import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field, replace

from lxml import etree

from ..ast import BlockQuote, Document, Footnote, Heading, Image, LineBreak, Node, Paragraph, Text
from ..images import ImageStore
from ..logger import log_warning
from ._docx_blocks import Flow, tidy
from ._office import box_preamble
from ._pptx_objects import read_chart, read_diagram, read_table, relationship_id
from ._pptx_text import TextReader, dominant_font, plain_text, single_line
from ._pptx_xml import DEFAULT_SECTION, R, Item, a, iter_items, p, slide_frame, xpath

_TITLE_TYPES = frozenset({"title", "ctrTitle"})
# Placeholders for the slide number, date, footer and header: the layout, not the content.
_CHROME_TYPES = frozenset({"dt", "ftr", "sldNum", "hdr"})
# Where a slide title can sit without a title placeholder, and how long it can be.
_TITLE_TOP = 0.25
_TITLE_WIDTH = 0.5
_TITLE_LENGTH = 150
_DEEPER_SLIDE = re.compile(r"^(#{1,4})[ \t]+")
_MEDIA = "p:nvPicPr/p:nvPr/*[local-name()='videoFile' or local-name()='audioFile' or local-name()='wavAudioFile']"


def parse_pptx(
    file_path: str,
    *,
    output_dir: str | None = None,
    image_dir_name: str | None = None,
    write_images: bool = True,
    keep_indentation: bool = False,
) -> Document:
    """Parse a PowerPoint presentation into an AST Document.

    Args:
        file_path: Path of the ``.pptx`` file.
        output_dir: Base directory for the extracted images (current directory if None).
        image_dir_name: Image subdirectory name inside ``output_dir``.
        write_images: Extract the pictures; when False they are left out.
        keep_indentation: Keep a paragraph's leading spaces, for code detection.

    Raises:
        ValueError: If the file is not a readable ``.pptx`` presentation.
    """
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError
    try:
        presentation = Presentation(file_path)
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, etree.XMLSyntaxError) as exc:
        raise ValueError(f"Not a readable PowerPoint (.pptx) presentation: {file_path} ({exc})") from exc
    return _Reader(presentation, output_dir, image_dir_name, write_images, keep_indentation).read()


@dataclass
class _Slide:
    """A slide read: its heading level and title, its content, and its speaker notes."""
    level: int
    title: list[Node]
    body: list[Node] = field(default_factory=list)
    notes: Footnote | None = None


def _placeholder_type(shape: etree._Element) -> str | None:
    """The type of the placeholder a shape is (``obj`` when it names none), None for a plain shape."""
    ph = shape.find(f"{p('nvSpPr')}/{p('nvPr')}/{p('ph')}")
    return None if ph is None else ph.get("type", "obj")


def _is_text_shape(shape: etree._Element) -> bool:
    return (shape.tag == p("sp") and shape.find(p("txBody")) is not None
            and bool("".join(t.text or "" for t in shape.iter(a("t"))).strip() or xpath(shape, ".//a14:m")))


class _Reader:
    """Reads one python-pptx presentation; see ``parse_pptx``."""

    def __init__(self, presentation, output_dir: str | None, image_dir_name: str | None, write_images: bool,
                 keep_indentation: bool = False):
        self.presentation = presentation
        self.width = float(presentation.slide_width)
        self.height = float(presentation.slide_height)
        self.write_images = write_images
        self.keep_indentation = keep_indentation
        self.dropped: Counter[str] = Counter()
        self._store = ImageStore(output_dir, image_dir_name)

    # ---- presentation -----------------------------------------------------

    def read(self) -> Document:
        children: list[Node] = []
        for name, slides in self._sections():
            heading = Heading(level=1, children=[Text(content=name)]) if name is not None else None
            if heading is not None:
                children.append(heading)
            for position, slide in enumerate(self._slide(s) for s in slides):
                if heading is not None and position == 0 and slide.level == 2 and plain_text(slide.title) == name:
                    # The section's own slide: its content belongs to the section heading.
                    heading.children.extend([slide.notes] if slide.notes else [])
                else:
                    notes = [slide.notes] if slide.notes else []
                    children.append(Heading(level=slide.level, children=[*slide.title, *notes]))
                children.extend(slide.body)
        props = self.presentation.core_properties
        doc = Document(children=children)
        if props.title:
            doc.title = [Text(content=props.title)]
        if props.author:
            doc.author = [Text(content=props.author)]
        if props.created:
            doc.date = [Text(content=props.created.date().isoformat())]
        if self.dropped:
            what = ", ".join(f"{name} ({count})" if count > 1 else name for name, count in self.dropped.items())
            log_warning(f"PowerPoint features without a counterpart were dropped: {what}.")
        return doc

    def _sections(self) -> list[tuple[str | None, list]]:
        """``(section name, slides)`` in order; a presentation with at most one section has no names."""
        slides = list(self.presentation.slides)
        sections = xpath(self.presentation.element, ".//p14:section")
        if len(sections) < 2:
            return [(None, slides)]
        by_id = {s.slide_id: s for s in slides}
        result, used = [], set()
        for section in sections:
            ids = [int(i) for i in xpath(section, "p14:sldIdLst/p14:sldId/@id")]
            result.append((section.get("name", ""), [by_id[i] for i in ids if i in by_id]))
            used.update(ids)
        if result[0][0] == DEFAULT_SECTION:  # the renderer's, for the slides before the first section
            result[0] = (None, result[0][1])
        rest = [s for s in slides if s.slide_id not in used]
        return result + [(None, rest)] if rest else result

    # ---- slides -----------------------------------------------------------

    def _slide(self, slide) -> _Slide:
        reader = TextReader(slide.part, self.dropped, keep_indentation=self.keep_indentation)
        items = list(iter_items(xpath(slide.element, "p:cSld/p:spTree")[0]))
        title = next((i for i in items if _placeholder_type(i.el) in _TITLE_TYPES), None) or self._guess_title(items)
        level, nodes = self._title(reader, title)
        main = self._main_shapes(items, title)
        result = _Slide(level=level, title=nodes, notes=self._notes(slide))
        for index, item in enumerate(items):
            if item is not title:
                result.body.extend(self._item(slide, reader, item, index, item.el in main))
        return result

    @staticmethod
    def _main_shapes(items: list[Item], title: Item | None) -> set[etree._Element]:
        """The text shapes that are the slide's body.

        Its content placeholders, even when empty; on a slide without any, the text box with the most text.
        """
        shapes = [i.el for i in items if i is not title and i.el.tag == p("sp")
                  and _placeholder_type(i.el) not in _CHROME_TYPES]
        placeholders = [el for el in shapes if _placeholder_type(el) is not None]
        if placeholders:
            return set(placeholders)
        texts = [el for el in shapes if _is_text_shape(el)]
        return {max(texts, key=lambda el: len("".join(t.text or "" for t in el.iter(a("t")))))} if texts else set()

    def _guess_title(self, items: list[Item]) -> Item | None:
        """The title of a slide without a title placeholder: a short, wide text box at the top."""
        candidates = []
        for item in items:
            if item.rect is None or item.el.tag != p("sp") or _placeholder_type(item.el) is not None:
                continue
            text = "".join(t.text or "" for t in item.el.iter(a("t"))).strip()
            if (text and len(text) <= _TITLE_LENGTH and len(item.el.findall(f".//{a('p')}")) <= 2
                    and item.rect.y <= _TITLE_TOP * self.height and item.rect.w >= _TITLE_WIDTH * self.width):
                sizes = [int(r.get("sz")) for r in item.el.iter(a("rPr")) if (r.get("sz") or "").isdigit()]
                candidates.append((item.rect.y, -max(sizes, default=0), len(candidates), item))
        return min(candidates)[3] if candidates else None

    @staticmethod
    def _title(reader: TextReader, title: Item | None) -> tuple[int, list[Node]]:
        """The heading level and the one-line title; a leading ``#`` run marks a slide below level 2."""
        nodes: list[Node] = []
        body = title.el.find(p("txBody")) if title is not None else None
        for paragraph in body.findall(a("p")) if body is not None else []:
            line, _ = reader.inline(paragraph)
            if line:
                nodes.extend([LineBreak()] if nodes else [])
                nodes.extend(line)
        # A title is already set apart; its bold text would be bold twice.
        nodes = [replace(n, bold=False) if isinstance(n, Text) else n for n in tidy(single_line(nodes))]
        if nodes and isinstance(nodes[0], Text):
            deeper = _DEEPER_SLIDE.match(nodes[0].content)
            if deeper:
                nodes[0] = replace(nodes[0], content=nodes[0].content[deeper.end():])
                return 2 + len(deeper.group(1)), nodes
        return 2, nodes

    def _notes(self, slide) -> Footnote | None:
        """The speaker notes as a footnote, one line per paragraph."""
        if not slide.has_notes_slide or slide.notes_slide.notes_placeholder is None:
            return None
        reader = TextReader(slide.notes_slide.part, self.dropped)
        content: list[Node] = []
        for paragraph in xpath(slide.notes_slide.notes_placeholder.element, "p:txBody/a:p"):
            nodes, _ = reader.inline(paragraph)
            if nodes:
                content.extend([LineBreak()] if content else [])
                content.extend(nodes)
        return Footnote(content=content) if content else None

    # ---- shapes -----------------------------------------------------------

    def _item(self, slide, reader: TextReader, item: Item, index: int, is_main: bool) -> list[Node]:
        el = item.el
        if el.tag == p("sp"):
            return self._text_shape(slide, reader, item, index, is_main)
        if el.tag == p("pic"):
            if xpath(el, _MEDIA) or xpath(el, "p:nvPicPr/p:nvPr/p:extLst//p14:media"):
                self.dropped["video or audio"] += 1
            return self._picture(slide, el, xpath(el, "p:nvPicPr/p:cNvPr"))
        if el.tag == p("graphicFrame"):
            return self._graphic(slide, reader, el)
        self.dropped["shape"] += 1
        return []

    def _text_shape(self, slide, reader: TextReader, item: Item, index: int, is_main: bool) -> list[Node]:
        el = item.el
        kind = _placeholder_type(el)
        if kind in _CHROME_TYPES:
            return []
        if not _is_text_shape(el):
            # An empty placeholder or text box is an artefact; an empty drawn shape is content lost.
            if kind is None and not xpath(el, "p:nvSpPr/p:cNvSpPr/@txBox"):
                self.dropped["shape"] += 1
            return []
        flow = Flow()
        reader.blocks(el.find(p("txBody")), flow, index, self._inherited_styles(slide, el) if is_main else [])
        flow.close()
        if is_main:
            return flow.out
        font = dominant_font(el.find(p("txBody")), xpath(self.presentation.element, "p:defaultTextStyle"))
        preamble = Paragraph(children=[Text(content=box_preamble(slide_frame(item.rect, self.width, self.height, font)))])
        return [BlockQuote(children=[preamble, *flow.out])]

    @staticmethod
    def _inherited_styles(slide, shape: etree._Element) -> list[etree._Element | None]:
        """The list styles a body placeholder inherits its bullets from: its layout's, then the master's."""
        kind = _placeholder_type(shape)
        if kind is None:
            return []
        idx = shape.find(f"{p('nvSpPr')}/{p('nvPr')}/{p('ph')}").get("idx", "0")
        layout = slide.slide_layout
        styles = xpath(layout.element, ".//p:sp[p:nvSpPr/p:nvPr/p:ph[@idx=$idx]]/p:txBody/a:lstStyle", idx=idx)
        if kind != "subTitle":
            styles.extend(xpath(layout.slide_master.element, "p:txStyles/p:bodyStyle"))
        return styles

    def _picture(self, slide, shape: etree._Element, props: list[etree._Element]) -> list[Node]:
        """The picture a shape holds as an image paragraph; nothing when it holds none or images are off."""
        blip = xpath(shape, ".//a:blip")
        embed = blip[0].get(f"{{{R}}}embed") if blip else None
        if not blip or embed is None:
            self.dropped["linked or missing image"] += 1
            return []
        if not self.write_images:
            return []
        try:
            part = slide.part.related_part(embed)
        except KeyError:
            self.dropped["linked or missing image"] += 1
            return []
        alt = (props[0].get("descr") or props[0].get("title") or "") if props else ""
        return [Paragraph(children=[Image(src=self._store.save(part.blob, part.partname.ext), alt=alt)])]

    def _graphic(self, slide, reader: TextReader, frame: etree._Element) -> list[Node]:
        """A table, a chart (as its data), a diagram (as its text) or an embedded object (as its preview)."""
        uri = (xpath(frame, "a:graphic/a:graphicData/@uri") or [""])[0]
        if uri.endswith("/table"):
            return [read_table(reader, xpath(frame, ".//a:tbl")[0])]
        if uri.endswith("/chart"):
            table = self._chart(slide, frame)
            if table is None:
                self.dropped["chart"] += 1
            return [table] if table is not None else []
        if uri.endswith("/diagram"):
            data = self._diagram(slide, frame)
            if data is None:
                self.dropped["diagram"] += 1
            return [data] if data is not None else []
        if xpath(frame, ".//a:blip"):
            return self._picture(slide, frame, xpath(frame, "p:nvGraphicFramePr/p:cNvPr"))
        self.dropped["embedded object"] += 1
        return []

    @staticmethod
    def _chart(slide, frame: etree._Element):
        rel_id = relationship_id(frame, ".//c:chart", "id")
        try:
            return read_chart(slide.part.related_part(rel_id).chart) if rel_id else None
        except (KeyError, ValueError):
            return None

    @staticmethod
    def _diagram(slide, frame: etree._Element):
        rel_id = relationship_id(frame, ".//dgm:relIds", "dm")
        try:
            return read_diagram(etree.fromstring(slide.part.related_part(rel_id).blob)) if rel_id else None
        except (KeyError, etree.XMLSyntaxError):
            return None
