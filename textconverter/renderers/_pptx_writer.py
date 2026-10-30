"""Text of a slide, written with python-pptx: lines (paragraphs, list items, code), runs, tables and notes.

Every paragraph carries its own bullet, indent and font size, so the result does not depend on the
template's master. Heights are estimated from the number of characters, which is enough to shrink the
text of a crowded slide until it fits.
"""
from __future__ import annotations

import math
from collections.abc import Callable
from dataclasses import dataclass, field

from lxml import etree

from ..ast import (
    Abstract, CodeBlock, CodeInline, Equation, Footnote, Image, LineBreak, Link, ListBlock, Node, Paragraph, Text,
)
from ..logger import log_warning
from ..parsers._omml import M_NS
from ..parsers._pptx_xml import A, MC, NS
from ._common import latex_reference_macro, plain_text, safe_href, text_runs
from ._omml import UnsupportedMath, latex_to_omml

CODE_FONT = "Consolas"
EMU_PER_PT = 12700
BODY_SIZE = 20.0
CODE_SIZE = 16.0
TABLE_SIZE = 14.0
MIN_SCALE = 0.5
_SPACE_AFTER = 6.0
_LINE_HEIGHT = 1.2
_INDENT = 342900
_BULLETS = ("•", "–", "▪")
_MATH_NS = (f'xmlns:mc="{MC}" xmlns:a="{A}" xmlns:a14="{NS["a14"]}" xmlns:m="{M_NS}"')


@dataclass
class Line:
    """One paragraph of a text frame, before it is written.

    ``bullet`` is ``(ordered, first number)`` for a list item; ``code`` is a line of a code block and
    ``display`` a display equation, each written in place of ``nodes``.
    """
    nodes: list[Node] = field(default_factory=list)
    level: int = 0
    bullet: tuple[bool, int] | None = None
    code: str | None = None
    display: Equation | None = None


def lines_of(blocks: list[Node]) -> list[Line]:
    """The paragraphs the blocks are written as: lists flattened with their levels, code split by line."""
    out: list[Line] = []
    for block in blocks:
        if isinstance(block, Paragraph):
            out.append(Line(nodes=block.children))
        elif isinstance(block, ListBlock):
            if out and out[-1].bullet is not None:
                out.append(Line())  # an empty paragraph keeps two lists in a row apart
            out.extend(_list_lines(block, 0))
        elif isinstance(block, CodeBlock):
            out.extend(Line(code=text) for text in block.code.split("\n"))
        elif isinstance(block, Equation):
            out.append(Line(display=block))
        elif isinstance(block, Abstract):
            out.extend(lines_of(block.children))
    return out


def _list_lines(block: ListBlock, level: int) -> list[Line]:
    out: list[Line] = []
    for item in block.items:
        rest = list(item.children)
        first = rest.pop(0) if rest and isinstance(rest[0], Paragraph) else Paragraph()
        nodes = list(first.children)
        if item.checked is not None:
            nodes.insert(0, Text(content="☑ " if item.checked else "☐ "))
        out.append(Line(nodes=nodes, level=level, bullet=(block.ordered, block.start)))
        for child in rest:
            if isinstance(child, ListBlock):
                out.extend(_list_lines(child, level + 1))
            else:
                out.extend(Line(level=level, nodes=line.nodes, code=line.code, display=line.display)
                           for line in lines_of([child]))
    return out


def line_size(line: Line) -> float:
    return CODE_SIZE if line.code is not None else max(12.0, BODY_SIZE - 2 * line.level)


def _wrapped(chars: int, size: float, width_pt: float, mono: bool) -> int:
    per_line = max(1.0, width_pt / (size * (0.6 if mono else 0.5)))
    return max(1, math.ceil(chars / per_line))


def lines_height(lines: list[Line], width_emu: float, scale: float) -> float:
    """The height in points the lines take at ``scale``."""
    total = 0.0
    for line in lines:
        size = line_size(line) * scale
        indent = (line.level + 1) * _INDENT if line.bullet else 0
        chars = len(line.code) if line.code is not None else len(plain_text(line.nodes))
        count = 2 if line.display else _wrapped(chars, size, (width_emu - indent) / EMU_PER_PT, line.code is not None)
        total += count * size * _LINE_HEIGHT + (0 if line.code is not None else _SPACE_AFTER * scale)
    return total


def cell_height(nodes: list[Node], width_emu: float, scale: float) -> float:
    """The height in points of a table cell holding ``nodes``, margins included."""
    size = TABLE_SIZE * scale
    chars = len(plain_text(nodes))
    return _wrapped(chars, size, width_emu / EMU_PER_PT - 14.4, False) * size * _LINE_HEIGHT + 14.4


def fit(measure: Callable[[float], float], available: float) -> float:
    """The font scale (between ``MIN_SCALE`` and 1) at which ``measure(scale)`` points fit in ``available``."""
    scale = 1.0
    for _ in range(5):
        height = measure(scale)
        if height <= available or scale <= MIN_SCALE:
            break
        scale = max(MIN_SCALE, scale * available / height)
    return scale


class TextWriter:
    """Writes lines and inline nodes into python-pptx text frames."""

    def fill(self, frame, lines: list[Line], scale: float, center: bool = False) -> None:
        """Replace the contents of a text frame with the lines."""
        from pptx.enum.text import PP_ALIGN
        frame.word_wrap = True
        for index, line in enumerate(lines or [Line()]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            if center:
                paragraph.alignment = PP_ALIGN.CENTER
            self._line(paragraph, line, scale)

    def _line(self, paragraph, line: Line, scale: float) -> None:
        size = line_size(line) * scale
        if line.code is None:
            from pptx.util import Pt
            paragraph.space_after = Pt(_SPACE_AFTER * scale)
        self._properties(paragraph, line)
        if line.code is not None:
            if line.code:
                self.run(paragraph, line.code, size, font=CODE_FONT)
            else:
                end = etree.SubElement(paragraph._p, f"{{{A}}}endParaRPr")
                etree.SubElement(end, f"{{{A}}}latin").set("typeface", CODE_FONT)
        elif line.display is not None:
            self.equation(paragraph, line.display, size, display=True)
        else:
            self.inline(paragraph, line.nodes, size)

    @staticmethod
    def _properties(paragraph, line: Line) -> None:
        """The bullet, indent and level of a paragraph, written out so no master supplies them."""
        props = paragraph._p.get_or_add_pPr()
        if line.bullet is None:
            props.set("marL", "0")
            props.set("indent", "0")
            etree.SubElement(props, f"{{{A}}}buNone")
            return
        ordered, start = line.bullet
        props.set("lvl", str(min(line.level, 8)))
        props.set("marL", str(_INDENT * (line.level + 1)))
        props.set("indent", str(-_INDENT * 3 // 4))
        if ordered:
            etree.SubElement(props, f"{{{A}}}buFont").set("typeface", "+mj-lt")
            number = etree.SubElement(props, f"{{{A}}}buAutoNum")
            number.set("type", "arabicPeriod")
            if start != 1:
                number.set("startAt", str(start))
        else:
            etree.SubElement(props, f"{{{A}}}buChar").set("char", _BULLETS[min(line.level, len(_BULLETS) - 1)])

    # ---- inline content -------------------------------------------------

    def inline(self, paragraph, nodes: list[Node], size: float, *, bold: bool = False, italic: bool = False,
               url: str | None = None) -> None:
        for node in text_runs(nodes):
            if isinstance(node, Text):
                for index, text in enumerate(node.content.split("\n")):
                    if index:
                        paragraph.add_line_break()
                    if text:
                        self.run(paragraph, text, size, bold=node.bold or bold, italic=node.italic or italic,
                                 strike=node.strikethrough, script=node.script, url=url)
            elif isinstance(node, CodeInline):
                self.run(paragraph, node.code, size, font=CODE_FONT, url=url)
            elif isinstance(node, LineBreak):
                paragraph.add_line_break()
            elif isinstance(node, Link):
                target = safe_href(node.url)
                self.inline(paragraph, node.content, size, bold=bold, italic=italic,
                            url=None if target.startswith("#") else target)
            elif isinstance(node, Equation):
                self.equation(paragraph, node, size, display=False)
            elif isinstance(node, (Image, Footnote)):
                continue
            else:
                macro = latex_reference_macro(node, plain_text)
                if macro:
                    self.run(paragraph, macro, size, url=url)

    @staticmethod
    def run(paragraph, text: str, size: float, *, bold: bool = False, italic: bool = False, strike: bool = False,
            script: str | None = None, font: str | None = None, url: str | None = None) -> None:
        from pptx.util import Pt
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold or None
        run.font.italic = italic or None
        props = run._r.get_or_add_rPr()
        if strike:
            props.set("strike", "sngStrike")
        if script:
            props.set("baseline", "30000" if script == "sup" else "-25000")
        if font:
            run.font.name = font
        if url:
            run.hyperlink.address = url

    @staticmethod
    def equation(paragraph, node: Equation, size: float, *, display: bool) -> None:
        """A native equation, or the LaTeX text when it cannot be translated."""
        try:
            math_element = latex_to_omml(node.code, node.env)
        except UnsupportedMath as exc:
            log_warning(f"Equation kept as LaTeX text, not translated to a PowerPoint equation ({exc}): "
                        f"{node.code[:60]}")
            TextWriter.run(paragraph, node.code, size, font=CODE_FONT)
            return
        wrapper = etree.fromstring(
            f'<mc:AlternateContent {_MATH_NS}><mc:Choice Requires="a14"><a14:m/></mc:Choice>'
            f'<mc:Fallback><a:r><a:t/></a:r></mc:Fallback></mc:AlternateContent>')
        holder = wrapper.find(f"{{{MC}}}Choice/{{{NS['a14']}}}m")
        if display:
            etree.SubElement(holder, f"{{{M_NS}}}oMathPara").append(math_element)
        else:
            holder.append(math_element)
        wrapper.find(f"{{{MC}}}Fallback/{{{A}}}r/{{{A}}}t").text = node.code
        paragraph._p.append(wrapper)

    def notes(self, frame, footnotes: list[Footnote]) -> None:
        """Speaker notes: one paragraph per line of each footnote."""
        first = True
        for note in footnotes:
            line: list[Node] = []
            for node in [*note.content, LineBreak()]:
                if isinstance(node, LineBreak):
                    if line:
                        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
                        first = False
                        self.inline(paragraph, line, 12.0)
                    line = []
                else:
                    line.append(node)
