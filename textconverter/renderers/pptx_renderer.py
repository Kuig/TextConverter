"""PowerPoint (.pptx) renderer: writes a Document AST with python-pptx onto a neutral 16:9 template.

A level 1 heading is a section and any deeper heading a slide (a level below 2 gets a ``#`` per level in its
title). The text of a slide goes in one content placeholder, shrunk until it fits; tables and pictures
follow it, and a quote headed by ``Text box (x .., y .., w .., h ..)`` becomes a text box at that frame.
Footnotes are the slide's speaker notes. The first slide uses the title layout when it looks like a title
and subtitle; a deck of bare level 1 headings, which would have no slide, gets a title slide for each.
"""
from __future__ import annotations

import io
import os
import uuid
from dataclasses import dataclass, field

from lxml import etree

from ..ast import (
    Abstract, BlockQuote, CodeBlock, Document, Equation, Footnote, Heading, Image, LineBreak, Link, ListBlock, Node,
    Paragraph, Table, Text, walk,
)
from ..logger import log_warning
from ..parsers._office import BoxFrame, parse_box_preamble
from ..parsers._pptx_xml import DEFAULT_SECTION, NS, P, a
from ._common import figure_image, lower_bibliography, plain_text, table_grid
from ._office import image_bytes, write_properties
from ._pptx_writer import (
    EMU_PER_PT, TABLE_SIZE, Line, TextWriter, cell_height, fit, line_size, lines_height, lines_of,
)

SLIDE_WIDTH = 12192000
SLIDE_HEIGHT = 6858000
_MARGIN = 457200
_TITLE_TOP = 250000
_TITLE_HEIGHT = 900000
_BODY_TOP = 1300000
_BODY_BOTTOM = SLIDE_HEIGHT - 350000
_GAP = 12 * EMU_PER_PT
_COLUMN_GAP = 300000
_PICTURE_SHARE = 0.4
_TITLE_SIZE = 32.0
_SECTIONS_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
_LAYOUT_TITLE, _LAYOUT_CONTENT, _LAYOUT_TITLE_ONLY = 0, 1, 5
# The longest paragraph under a level 2 heading that still reads as a subtitle.
_SUBTITLE_LENGTH = 80


def render_pptx(doc: Document, output_path: str, base_dir: str | None = None) -> None:
    """Write the document as a PowerPoint file.

    Args:
        doc: The document to render.
        output_path: Where to write the ``.pptx`` file.
        base_dir: Directory that relative local image paths resolve against
            (the current directory when None).
    """
    from pptx import Presentation
    _Deck(Presentation(), base_dir or os.getcwd()).write(doc, output_path)


@dataclass
class _Picture:
    data: bytes
    alt: str
    size: tuple[int, int]


@dataclass
class _Plan:
    """What one slide will hold, sorted out of the document's blocks."""
    level: int
    title: list[Node]
    section: int | None
    blocks: list[Node] = field(default_factory=list)
    quotes: list[BlockQuote] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    pictures: list[Image] = field(default_factory=list)
    registered: bool = False
    subheading: bool = False  # made by a heading deeper than level 1
    bare_section: bool = False  # a level 1 heading alone, in a deck that would otherwise have no slide

    def footnotes(self) -> list[Footnote]:
        nodes = [*self.title, *self.blocks, *self.quotes, *self.tables]
        return [n for node in nodes for n in walk(node) if isinstance(n, Footnote)]


class _Deck:
    def __init__(self, presentation, base_dir: str):
        self.prs = presentation
        self.base_dir = base_dir
        self.text = TextWriter()
        self._widen_layouts()

    def _widen_layouts(self) -> None:
        """Make the template 16:9: the master and layouts are drawn for 4:3, so their frames are stretched."""
        factor = SLIDE_WIDTH / self.prs.slide_width
        self.prs.slide_width, self.prs.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT
        master = self.prs.slide_master
        for part in (master, *master.slide_layouts):
            for off in part.element.iter(a("off")):
                off.set("x", str(round(int(off.get("x")) * factor)))
            for ext in part.element.iter(a("ext")):
                if ext.get("cx"):
                    ext.set("cx", str(round(int(ext.get("cx")) * factor)))

    def write(self, doc: Document, output_path: str) -> None:
        write_properties(self.prs.core_properties, doc)
        lower_bibliography(doc)
        plans, sections = self._plan(doc)
        slides = []
        for index, plan in enumerate(plans):
            slides.append(self._slide(plan, first=index == 0))
        if sections:
            self._write_sections(plans, slides, sections)
        self.prs.save(output_path)

    # ---- the document's blocks, sorted into slides ----------------------

    def _plan(self, doc: Document) -> tuple[list[_Plan], list[str]]:
        plans: list[_Plan] = []
        sections: list[str] = []
        section_plans: list[_Plan] = []
        current: _Plan | None = None

        def register(plan: _Plan) -> _Plan:
            if not plan.registered:
                plan.registered = True
                plans.append(plan)
            return plan

        for node in doc.children:
            if isinstance(node, Heading):
                if node.level == 1:
                    sections.append(plain_text(node.children).strip())
                current = _Plan(level=2 if node.level == 1 else node.level, title=node.children,
                                section=len(sections) - 1 if sections else None, subheading=node.level != 1)
                if node.level == 1:
                    section_plans.append(current)
                if node.level != 1 or any(isinstance(n, Footnote) for n in node.children):
                    register(current)
            else:
                if current is None:
                    current = _Plan(level=2, title=[], section=None)
                self._add(register(current), node)
        if not plans:
            for plan in section_plans:
                plan.bare_section = True
                register(plan)
        return plans, sections

    def _add(self, plan: _Plan, node: Node) -> None:
        if isinstance(node, Paragraph):
            rest = []
            for child in node.children:
                image = child.content[0] if isinstance(child, Link) and len(child.content) == 1 else child
                if isinstance(image, Image):
                    plan.pictures.append(image)
                else:
                    rest.append(child)
            if any(not (isinstance(n, Text) and not n.content.strip()) for n in rest):
                plan.blocks.append(Paragraph(children=rest))
        elif isinstance(node, BlockQuote):
            plan.quotes.append(node)
        elif isinstance(node, Table):
            plan.tables.append(node)
        elif isinstance(node, Abstract):
            for child in node.children:
                self._add(plan, child)
        elif isinstance(node, (ListBlock, CodeBlock, Equation)):
            plan.blocks.append(node)

    # ---- one slide --------------------------------------------------------

    def _slide(self, plan: _Plan, first: bool):
        pictures = self._load_pictures(plan)
        captions = [Paragraph(children=[Text(content=plain_text(c.caption), italic=True)])
                    for c in [*plan.tables, *plan.pictures] if c.caption]
        lines = lines_of([*plan.blocks, *captions])
        title_layout = _reads_as_title(plan, pictures, first)
        layout = _LAYOUT_TITLE if title_layout else _LAYOUT_CONTENT if lines else _LAYOUT_TITLE_ONLY
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout])
        self._title(slide, plan, centered=title_layout)
        if title_layout:
            self._subtitle(slide, lines)
        else:
            self._body(slide, plan, lines, pictures)
        notes = plan.footnotes()
        if notes:
            self.text.notes(slide.notes_slide.notes_text_frame, notes)
        return slide

    def _title(self, slide, plan: _Plan, centered: bool) -> None:
        from pptx.enum.text import PP_ALIGN
        shape = slide.shapes.title
        width = SLIDE_WIDTH - 2 * _MARGIN
        if centered:
            shape.left, shape.top, shape.width, shape.height = _MARGIN, 1900000, width, 1400000
        else:
            shape.left, shape.top, shape.width, shape.height = _MARGIN, _TITLE_TOP, width, _TITLE_HEIGHT
        nodes = plan.title
        if plan.level > 2:
            nodes = [Text(content="#" * (plan.level - 2) + " "), *nodes]
        size = fit_title(plain_text([n for n in nodes if not isinstance(n, Footnote)]), width,
                         40.0 if centered else _TITLE_SIZE)
        frame = shape.text_frame
        frame.word_wrap = True
        if centered:
            frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        self.text.inline(frame.paragraphs[0], nodes, size)

    def _subtitle(self, slide, lines: list[Line]) -> None:
        shape = slide.placeholders[1]
        if not lines:
            shape._element.getparent().remove(shape._element)
            return
        shape.left, shape.top, shape.width, shape.height = _MARGIN * 2, 3500000, SLIDE_WIDTH - 4 * _MARGIN, 1800000
        self.text.fill(shape.text_frame, lines, fit(lambda s: lines_height(lines, shape.width, s), 1800000 / EMU_PER_PT),
                       center=True)

    def _body(self, slide, plan: _Plan, lines: list[Line], pictures: list[_Picture]) -> None:
        body_width = SLIDE_WIDTH - 2 * _MARGIN
        has_left = bool(lines or plan.tables)
        left_width = body_width * (1 - _PICTURE_SHARE) - _COLUMN_GAP if pictures and has_left else body_width
        placed, loose = [], []
        for quote in plan.quotes:
            frame = parse_box_preamble(plain_text(quote.children[0].children)) if _starts_with_paragraph(quote) else None
            (placed if frame is not None and frame.placed else loose).append((quote, frame))
        loose_lines = [lines_of(_quote_blocks(q, bool(f))) for q, f in loose]
        grids = [table_grid(t) for t in plan.tables]

        def measure(scale: float) -> float:
            total = lines_height(lines, left_width, scale) if lines else 0.0
            for headers, rows, columns in grids:
                total += self._table_height(headers, rows, columns, left_width, scale)
            total += sum(lines_height(block, left_width, scale) for block in loose_lines)
            return total + _GAP / EMU_PER_PT * max(0, len(grids) + len(loose_lines))

        available = (_BODY_BOTTOM - _BODY_TOP) / EMU_PER_PT
        scale = fit(measure, available)
        y = _BODY_TOP
        if lines:
            shape = slide.placeholders[1]
            height = int(lines_height(lines, left_width, scale) * EMU_PER_PT)
            shape.left, shape.top, shape.width, shape.height = _MARGIN, y, int(left_width), height
            self.text.fill(shape.text_frame, lines, scale)
            y += height + _GAP
        for table, (headers, rows, columns) in zip(plan.tables, grids):
            y += self._table(slide, table, headers, rows, columns, _MARGIN, y, left_width, scale) + _GAP
        for block in loose_lines:
            height = int(lines_height(block, left_width, scale) * EMU_PER_PT)
            self._text_box(slide, block, _MARGIN, y, int(left_width), height, scale)
            y += height + _GAP
        for quote, frame in placed:
            self._placed_quote(slide, quote, frame)
        self._pictures(slide, pictures, body_width, left_width if has_left else 0)

    def _text_box(self, slide, lines: list[Line], x: int, y: int, width: int, height: int, scale: float) -> None:
        box = slide.shapes.add_textbox(x, y, width, height)
        self.text.fill(box.text_frame, lines, scale)

    def _placed_quote(self, slide, quote: BlockQuote, frame: BoxFrame) -> None:
        x, y, w, h = (round(v / 100 * total) for v, total in zip(
            (frame.x, frame.y, frame.w, frame.h), (SLIDE_WIDTH, SLIDE_HEIGHT) * 2))
        lines = lines_of(_quote_blocks(quote, True))
        start = frame.font / line_size(lines[0]) if frame.font and lines else 1.0
        scale = fit(lambda s: lines_height(lines, w, s), max(h, 300000) / EMU_PER_PT, start)
        height = max(h, int(lines_height(lines, w, scale) * EMU_PER_PT))
        self._text_box(slide, lines, x, y, max(w, 300000), height, scale)

    # ---- tables and pictures ----------------------------------------------

    @staticmethod
    def _table_height(headers, rows, columns: int, width: float, scale: float) -> float:
        if not columns:
            return 0.0
        cell_width = width / columns
        return sum(max(cell_height(c.children, cell_width, scale) for c in row)
                   for row in ([headers] if headers else []) + rows)

    def _table(self, slide, table: Table, headers, rows, columns: int, x: int, y: int, width: float,
               scale: float) -> int:
        grid = ([headers] if headers else []) + rows
        if not columns or not grid:
            return 0
        height = int(self._table_height(headers, rows, columns, width, scale) * EMU_PER_PT)
        shape = slide.shapes.add_table(len(grid), columns, x, y, int(width), height)
        shape.table.first_row = bool(headers)
        cell_width = width / columns
        for r, row in enumerate(grid):
            shape.table.rows[r].height = int(max(cell_height(c.children, cell_width, scale) for c in row) * EMU_PER_PT)
            for c, cell in enumerate(row):
                self.text.inline(shape.table.cell(r, c).text_frame.paragraphs[0], cell.children, TABLE_SIZE * scale)
        return height

    def _load_pictures(self, plan: _Plan) -> list[_Picture]:
        """The slide's pictures read and checked; one that cannot be embedded becomes a line of text."""
        from PIL import Image as PilImage
        loaded = []
        for image in plan.pictures:
            blob = image_bytes(image.src, self.base_dir)
            try:
                if blob is None:
                    raise ValueError("not found")
                with PilImage.open(io.BytesIO(blob)) as picture:
                    size = picture.size
                    if picture.format not in ("PNG", "JPEG", "GIF", "BMP", "TIFF"):
                        out = io.BytesIO()
                        picture.convert("RGBA" if picture.mode in ("RGBA", "LA", "P") else "RGB").save(out, "PNG")
                        blob = out.getvalue()
                loaded.append(_Picture(data=blob, alt=image.alt or "", size=size))
            except (ValueError, OSError) as exc:
                log_warning(f"Image not embedded in the PowerPoint file ({exc}): {image.src[:80]}")
                fallback = image.alt or ("image" if image.src.startswith("data:") else image.src)
                plan.blocks.append(Paragraph(children=[Text(content=f"[{fallback}]")]))
        return loaded

    @staticmethod
    def _pictures(slide, pictures: list[_Picture], body_width: float, left_width: float) -> None:
        if not pictures:
            return
        x = _MARGIN + (left_width + _COLUMN_GAP if left_width else 0)
        width = body_width - (left_width + _COLUMN_GAP if left_width else 0)
        height = (_BODY_BOTTOM - _BODY_TOP - _GAP * (len(pictures) - 1)) / len(pictures)
        y = _BODY_TOP
        for picture in pictures:
            ratio = min(width / picture.size[0], height / picture.size[1])
            w, h = int(picture.size[0] * ratio), int(picture.size[1] * ratio)
            shape = slide.shapes.add_picture(io.BytesIO(picture.data), int(x + (width - w) / 2), int(y), w, h)
            shape._element.nvPicPr.cNvPr.set("descr", picture.alt)
            y += h + _GAP

    # ---- sections -------------------------------------------------------------

    def _write_sections(self, plans: list[_Plan], slides: list, sections: list[str]) -> None:
        """Group the slides in ``p14:section`` elements; slides before the first section get a default one."""
        members: dict[int | None, list[int]] = {}
        for plan, slide in zip(plans, slides):
            members.setdefault(plan.section, []).append(slide.slide_id)
        groups = ([(DEFAULT_SECTION, members.get(None, []))] if members.get(None) else [])
        groups += [(name, members.get(i, [])) for i, name in enumerate(sections)]
        root = self.prs.element
        ext_list = root.find(f"{{{P}}}extLst")
        if ext_list is None:
            ext_list = etree.SubElement(root, f"{{{P}}}extLst")
        ext = etree.SubElement(ext_list, f"{{{P}}}ext", uri=_SECTIONS_URI)
        listing = etree.SubElement(ext, f"{{{NS['p14']}}}sectionLst", nsmap={"p14": NS["p14"]})
        for name, ids in groups:
            section = etree.SubElement(listing, f"{{{NS['p14']}}}section", name=name, id="{%s}" % str(uuid.uuid4()).upper())
            slide_ids = etree.SubElement(section, f"{{{NS['p14']}}}sldIdLst")
            for slide_id in ids:
                etree.SubElement(slide_ids, f"{{{NS['p14']}}}sldId", id=str(slide_id))


def _reads_as_title(plan: _Plan, pictures: list[_Picture], first: bool) -> bool:
    """Whether a slide takes the title layout: a bare section, or a first slide that is a title and subtitle.

    A level 2 slide qualifies with one short line; a section's own text or
    text before any heading with one or two paragraphs.
    """
    if plan.bare_section:
        return True
    if not first or plan.level != 2 or plan.quotes or plan.tables or pictures:
        return False
    if plan.subheading:
        return len(plan.blocks) == 1 and _is_short_line(plan.blocks[0])
    return 0 < len(plan.blocks) <= 2 and all(isinstance(b, Paragraph) for b in plan.blocks)


def _is_short_line(block: Node) -> bool:
    if not isinstance(block, Paragraph) or any(isinstance(n, LineBreak) for n in walk(block)):
        return False
    text = plain_text(block.children).strip()
    return "\n" not in text and len(text) <= _SUBTITLE_LENGTH


def fit_title(text: str, width_emu: float, size: float) -> float:
    """The title's font size: ``size``, reduced (to 18 at least) until the title fits in two lines."""
    per_line = width_emu / EMU_PER_PT / 0.55
    while size > 18.0 and len(text) * size > 2 * per_line:
        size -= 2
    return size


def _starts_with_paragraph(quote: BlockQuote) -> bool:
    return bool(quote.children) and isinstance(quote.children[0], Paragraph)


def _quote_blocks(quote: BlockQuote, placed: bool) -> list[Node]:
    """The blocks a quote writes: all of its children, without the frame line when it had one."""
    return quote.children[1:] if placed else quote.children
