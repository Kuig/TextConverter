"""PowerPoint (.pptx) input: the real fixture Test_slides.pptx and presentations built with python-pptx.

Usage: python -m unittest discover -s tests -p "test_pptx.py"
"""
from __future__ import annotations
import io
import os
import tempfile
import zipfile
from unittest.mock import patch

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt

import textconverter
from textconverter.api import convert
from textconverter.ast import (
    BlockQuote, CodeBlock, CodeInline, Document, Equation, Footnote, Heading, Image, LineBreak, Link, ListBlock,
    Paragraph, Table, Text, walk,
)
from textconverter.formats import infer_input_format, infer_input_format_from_media_type
from textconverter.parsers._office import parse_box_preamble
from textconverter.parsers._pptx_xml import Rect
from textconverter.parsers.pptx_parser import parse_pptx
from textconverter.renderers._common import plain_text
from _support import BaseConverterTest, setUpModule, tearDownModule, INPUT_DIR

FIXTURE = INPUT_DIR / "Test_slides.pptx"
_NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
       'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" '
       'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
       'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
       'xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"')


def _of(doc, cls):
    return [n for n in walk(doc) if isinstance(n, cls)]


def _headings(doc) -> list[tuple[int, str]]:
    return [(h.level, plain_text(h.children)) for h in _of(doc, Heading)]


def _png() -> io.BytesIO:
    from PIL import Image as PilImage
    buffer = io.BytesIO()
    PilImage.new("RGB", (4, 3), "red").save(buffer, "PNG")
    buffer.seek(0)
    return buffer


class PptxFixtureTest(BaseConverterTest):
    """Parses Test_slides.pptx once; the tests read the same AST."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.tmp.cleanup)
        with patch("textconverter.parsers.pptx_parser.log_warning") as warn:
            cls.doc = parse_pptx(str(FIXTURE), output_dir=cls.tmp.name, image_dir_name="img")
        cls.warnings = [call.args[0] for call in warn.call_args_list]


class TestFixtureStructure(PptxFixtureTest):

    def test_metadata_comes_from_the_file_properties(self):
        self.assertEqual(plain_text(self.doc.title), "Titolo dai metadati")
        self.assertEqual(plain_text(self.doc.author), "Giorgio")

    def test_sections_are_level_one_and_slides_level_two(self):
        headings = _headings(self.doc)
        self.assertEqual([t for level, t in headings if level == 1],
                         ["Sezione predefinita", "Sezione 1", "Sezione 2"])
        self.assertEqual(sum(1 for level, _ in headings if level == 2), 37)
        self.assertEqual({level for level, _ in headings}, {1, 2})

    def test_slide_titles_are_one_line_and_not_bold(self):
        for heading in _of(self.doc, Heading):
            for node in heading.children:
                if isinstance(node, Text):
                    self.assertFalse(node.bold)
                    self.assertNotIn("\n", node.content)
        self.assertIn((2, "Giorgio Presti Researcher @ LIM"), _headings(self.doc))

    def test_speaker_notes_are_footnotes_on_the_heading(self):
        notes = [h for h in _of(self.doc, Heading) if any(isinstance(n, Footnote) for n in h.children)]
        self.assertEqual(len(notes), 2)
        self.assertIn("ambiente (environment)", plain_text(next(
            n for n in notes[0].children if isinstance(n, Footnote)).content))

    def test_secondary_text_boxes_are_quotes_with_their_frame(self):
        code = next(q for q in _of(self.doc, BlockQuote) if _of(q, CodeBlock))
        frame = parse_box_preamble(plain_text(code.children[0].children))
        self.assertEqual((frame.x, frame.y, frame.w, frame.h), (20, 62, 41, 15))
        self.assertIn("export LAVORO", _of(code, CodeBlock)[0].code)

    def test_every_text_box_of_the_fixture_tells_its_font(self):
        frames = [parse_box_preamble(plain_text(q.children[0].children)) for q in _of(self.doc, BlockQuote)
                  if isinstance(q.children[0], Paragraph)]
        frames = [f for f in frames if f is not None]
        self.assertTrue(frames)
        self.assertTrue(all(f.font for f in frames), [f for f in frames if not f.font])

    def test_lists_nest_and_tables_have_headers(self):
        self.assertTrue(any(_of(item, ListBlock) for lst in _of(self.doc, ListBlock) for item in lst.items))
        tables = _of(self.doc, Table)
        self.assertEqual(len(tables), 2)
        self.assertTrue(all(t.headers for t in tables))

    def test_pictures_are_written_once_per_content(self):
        sources = {i.src for i in _of(self.doc, Image)}
        self.assertTrue(sources)
        for src in sources:
            self.assertTrue(os.path.exists(os.path.join(self.tmp.name, src)), src)

    def test_one_warning_lists_what_was_dropped(self):
        self.assertEqual(len(self.warnings), 1)
        self.assertIn("shape", self.warnings[0])


class TestGeneratedPresentations(BaseConverterTest):

    def setUp(self):
        super().setUp()
        self.prs = Presentation()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def _parse(self, **kwargs):
        path = os.path.join(self.tmp.name, "deck.pptx")
        self.prs.save(path)
        with patch("textconverter.parsers.pptx_parser.log_warning") as warn:
            doc = parse_pptx(path, output_dir=self.tmp.name, image_dir_name="img", **kwargs)
        self.warnings = [call.args[0] for call in warn.call_args_list]
        return doc

    def _content_slide(self, title: str):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        slide.shapes.title.text = title
        return slide

    def test_hashes_in_a_title_make_a_deeper_slide(self):
        self._content_slide("Plain")
        self._content_slide("# Deeper")
        self._content_slide("## Deepest")
        self.assertEqual(_headings(self._parse()), [(2, "Plain"), (3, "Deeper"), (4, "Deepest")])

    def test_bullets_levels_and_links(self):
        slide = self._content_slide("Lists")
        frame = slide.placeholders[1].text_frame
        frame.text = "first"
        frame.add_paragraph().text = "nested"
        frame.paragraphs[1].level = 1
        run = frame.add_paragraph().add_run()
        run.text = "site"
        run.hyperlink.address = "https://example.com"
        doc = self._parse()
        outer = _of(doc, ListBlock)[0]
        self.assertEqual(len(outer.items), 2)
        self.assertEqual(len(_of(outer.items[0], ListBlock)), 1)
        self.assertEqual(_of(doc, Link)[0].url, "https://example.com")

    def test_text_boxes_outside_a_placeholder_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(8), Inches(0.8))
        title.text_frame.text = "Guessed title"
        body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        body.text_frame.text = "The long main text of the slide"
        slide.shapes.add_textbox(Inches(8), Inches(5), Inches(1), Inches(1)).text_frame.text = "side"
        doc = self._parse()
        self.assertEqual(_headings(doc), [(2, "Guessed title")])
        quotes = _of(doc, BlockQuote)
        self.assertEqual(len(quotes), 1)
        self.assertIsNotNone(parse_box_preamble(plain_text(quotes[0].children[0].children)))
        self.assertIn("side", plain_text(quotes[0].children[1].children))

    def test_slide_without_a_title_has_an_empty_heading(self):
        self.prs.slides.add_slide(self.prs.slide_layouts[6]).shapes.add_textbox(
            Inches(1), Inches(4), Inches(2), Inches(1)).text_frame.text = "low"
        self.assertEqual(_headings(self._parse()), [(2, "")])

    def test_chart_becomes_a_table(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        slide.shapes.title.text = "Data"
        data = CategoryChartData()
        data.categories = ["A", "B"]
        data.add_series("S1", (1, 2.5))
        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), data).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = "Sales"
        table = _of(self._parse(), Table)[0]
        self.assertEqual(plain_text(table.caption), "Sales")
        self.assertEqual([plain_text(c.children) for c in table.headers], ["", "S1"])
        self.assertEqual([[plain_text(c.children) for c in r.cells] for r in table.rows],
                         [["A", "1"], ["B", "2.5"]])

    def test_hidden_slides_are_kept(self):
        slide = self._content_slide("Hidden")
        slide.element.set("show", "0")
        self.assertEqual(_headings(self._parse()), [(2, "Hidden")])

    def test_slide_titled_like_its_section_merges_into_the_section(self):
        first = self._content_slide("Intro")
        first.placeholders[1].text_frame.text = "welcome"
        self._content_slide("Next")
        ids = [s.slide_id for s in self.prs.slides]
        sections = "".join(
            f'<p14:section name="{name}" id="{{00000000-0000-0000-0000-00000000000{i}}}"><p14:sldIdLst>'
            f'<p14:sldId id="{sid}"/></p14:sldIdLst></p14:section>'
            for i, (name, sid) in enumerate(zip(("Intro", "Part two"), ids)))
        ext = etree.fromstring(
            f'<p:extLst {_NS}><p:ext uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}"><p14:sectionLst>'
            f'{sections}</p14:sectionLst></p:ext></p:extLst>')
        self.prs.element.append(ext)
        doc = self._parse()
        self.assertEqual(_headings(doc), [(1, "Intro"), (1, "Part two"), (2, "Next")])
        self.assertIsInstance(doc.children[1], ListBlock)

    def test_smartart_text_becomes_a_nested_list(self):
        slide = self._content_slide("Diagram")
        model = (f'<dgm:dataModel {_NS}><dgm:ptLst>'
                 '<dgm:pt modelId="0" type="doc"/>'
                 '<dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>Root</a:t></a:r></a:p></dgm:t></dgm:pt>'
                 '<dgm:pt modelId="2"><dgm:t><a:p><a:r><a:t>Leaf</a:t></a:r></a:p></dgm:t></dgm:pt>'
                 '</dgm:ptLst><dgm:cxnLst>'
                 '<dgm:cxn modelId="10" srcId="0" destId="1" srcOrd="0"/>'
                 '<dgm:cxn modelId="11" srcId="1" destId="2" srcOrd="0"/>'
                 '</dgm:cxnLst></dgm:dataModel>')
        part = Part(PackURI("/ppt/diagrams/data1.xml"),
                    "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
                    self.prs.part.package, model.encode())
        rel_id = slide.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData")
        frame = etree.fromstring(
            f'<p:graphicFrame {_NS}><p:nvGraphicFramePr><p:cNvPr id="9" name="Diagram"/><p:cNvGraphicFramePr/>'
            '<p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="0" y="0"/><a:ext cx="100" cy="100"/></p:xfrm>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
            f'<dgm:relIds r:dm="{rel_id}" r:lo="rId1" r:qs="rId1" r:cs="rId1"/></a:graphicData></a:graphic>'
            '</p:graphicFrame>')
        slide.shapes._spTree.append(frame)
        lists = _of(self._parse(), ListBlock)
        self.assertEqual(plain_text(lists[0].items[0].children[0].children), "Root")
        self.assertEqual(plain_text(lists[0].items[0].children[1].items[0].children[0].children), "Leaf")

    def test_equation_and_media_and_missing_pictures(self):
        slide = self._content_slide("Math")
        paragraph = slide.placeholders[1].text_frame.paragraphs[0]._p
        paragraph.append(etree.fromstring(
            f'<mc:AlternateContent {_NS}><mc:Choice Requires="a14"><a14:m><m:oMathPara><m:oMath>'
            '<m:r><m:t>x</m:t></m:r></m:oMath></m:oMathPara></a14:m></mc:Choice></mc:AlternateContent>'))
        picture = slide.shapes.add_picture(_png(), 0, 0)
        picture._element.nvPicPr.nvPr.append(etree.fromstring(
            f'<a:videoFile {_NS} r:link="rId99"/>'))
        doc = self._parse()
        self.assertEqual(_of(doc, Equation)[0].code, "x")
        self.assertEqual(len(_of(doc, Image)), 1)
        self.assertIn("video or audio", self.warnings[0])

    def test_images_can_be_left_out(self):
        self._content_slide("Pic").shapes.add_picture(_png(), 0, 0)
        self.assertEqual(_of(self._parse(write_images=False), Image), [])
        self.assertFalse(os.path.exists(os.path.join(self.tmp.name, "img")))

    def test_not_a_presentation_is_a_value_error(self):
        path = os.path.join(self.tmp.name, "bad.pptx")
        with open(path, "wb") as f:
            f.write(b"not a zip")
        with self.assertRaisesRegex(ValueError, "Not a readable PowerPoint"):
            parse_pptx(path)


class TestBoxPreamble(BaseConverterTest):

    def _frame(self, text: str):
        from textconverter.parsers._office import parse_box_preamble as parse
        return parse(text)

    def test_every_form_reads_and_writes_back(self):
        from textconverter.parsers._office import BoxFrame, box_preamble as write
        for text, frame in (
            ("Text box (x 10%, y 10%, w 30%, h 20%)", BoxFrame(x=10, y=10, w=30, h=20)),
            ("Text box (x 10%, y 50%, w 30%, h 20%, font 18pt)", BoxFrame(x=10, y=50, w=30, h=20, font=18)),
            ("Text box (x -5%, y 10%, w 30%, h 20%, font 10.5pt)", BoxFrame(x=-5, y=10, w=30, h=20, font=10.5)),
            ("Text box (x 35%, w 23%, h 5%, font 11pt)", BoxFrame(x=35, w=23, h=5, font=11)),
            ("Text box (C4:M13, font 11pt)", BoxFrame(cells="C4:M13", font=11)),
            ("Text box (C4:M13)", BoxFrame(cells="C4:M13")),
            ("Text box (font 11pt)", BoxFrame(font=11)),
        ):
            with self.subTest(text):
                self.assertEqual(self._frame(text), frame)
                self.assertEqual(write(frame), text)

    def test_what_is_not_a_preamble(self):
        for text in ("An ordinary quote", "Text box", "Text box ()", "Text box (hello)", "A Text box (x 1%)"):
            with self.subTest(text):
                self.assertIsNone(self._frame(text))

    def test_a_frame_from_a_rectangle(self):
        from textconverter.parsers._office import box_preamble as write
        from textconverter.parsers._pptx_xml import slide_frame
        frame = slide_frame(Rect(1219200, 685800, 3657600, 1371600), 12192000, 6858000, 24)
        self.assertEqual(write(frame), "Text box (x 10%, y 10%, w 30%, h 20%, font 24pt)")


class TestCodeDetection(BaseConverterTest):
    CODE = "def f(x):\n    y = x + 1\n    return y"

    def _markdown(self, **options) -> str:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(100, 100, 3000000, 3000000)
        box.text_frame.text = "Some introduction to the slide with enough plain words in it"
        other = slide.shapes.add_textbox(100, 3200000, 3000000, 3000000)
        lines = self.CODE.split("\n")
        other.text_frame.text = lines[0]
        for line in lines[1:]:
            other.text_frame.add_paragraph().text = line
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "code.pptx")
            prs.save(path)
            return convert(path, "markdown", textconverter.ConvertOptions(image_handling="discard", **options))

    def test_code_in_a_text_box_keeps_the_frame_line(self):
        markdown = self._markdown(code_parsing=True)
        self.assertIn("```python", markdown)
        self.assertIn("```python" + chr(10) + self.CODE + chr(10) + "```", markdown.replace("> ", ""))
        self.assertRegex(markdown, r"> Text box .x \d+%")

    def test_nothing_changes_without_the_flag(self):
        self.assertNotIn("```", self._markdown())

    def test_the_short_text_of_real_slides_is_not_code(self):
        markdown = convert(str(INPUT_DIR / "Test_slides.pptx"), "markdown",
                           textconverter.ConvertOptions(image_handling="discard", code_parsing=True))
        self.assertEqual(markdown, convert(str(INPUT_DIR / "Test_slides.pptx"), "markdown",
                                           textconverter.ConvertOptions(image_handling="discard")))


class TestPptxThroughTheApi(BaseConverterTest):

    def test_markdown_from_the_fixture(self):
        with tempfile.TemporaryDirectory() as tmp:
            text = convert(str(FIXTURE), "markdown", output_dir=tmp, image_dir_name="images")
        self.assertIn("# Sezione 1", text)
        self.assertIn("![", text)

    def test_formats_are_inferred(self):
        self.assertEqual(infer_input_format("a.PPTX"), "pptx")
        self.assertEqual(infer_input_format_from_media_type(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"), "pptx")


class PptxOutputTest(BaseConverterTest):
    """Helpers for tests that write a PowerPoint file."""

    def _save(self, source: str, **options) -> str:
        """Convert Markdown ``source`` to a .pptx in a temporary folder; return the .pptx path."""
        from textconverter import ConvertOptions, save_to_file
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        src = os.path.join(tmp.name, "in.md")
        with open(src, "w", encoding="utf-8") as f:
            f.write(source)
        out = os.path.join(tmp.name, "out.pptx")
        save_to_file(src, out, ConvertOptions(**options))
        self.assertEqual(sorted(os.listdir(tmp.name)), ["in.md", "out.pptx"], "no image folder or temp file left")
        return out

    def _read(self, path: str):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        return parse_pptx(path, output_dir=tmp.name, image_dir_name="img")

    def _roundtrip(self, source: str):
        return self._read(self._save(source))

    def _render(self, doc: Document) -> str:
        """Render ``doc`` straight to a .pptx under tests/Output and return its path."""
        from textconverter.renderers.pptx_renderer import render_pptx
        path = str(self._out("render_test.pptx"))
        render_pptx(doc, path)
        return path


class TestPptxRenderer(PptxOutputTest):

    def test_headings_sections_lists_and_text_survive(self):
        source = ("# Part A\n\n## One\n\nSome text\n\n- a\n    - b\n- c\n\n### Sub\n\n3. x\n4. y\n\n"
                  "# Part B\n\n## Two\n\n```\ncode line\n\nlast\n```\n")
        doc = self._roundtrip(source)
        self.assertEqual(_headings(doc), [(1, "Part A"), (2, "One"), (3, "Sub"), (1, "Part B"), (2, "Two")])
        lists = _of(doc, ListBlock)
        self.assertEqual([plain_text(i.children[0].children) for i in lists[0].items], ["a", "c"])
        self.assertEqual(len(_of(lists[0].items[0], ListBlock)), 1)
        self.assertEqual((lists[-1].ordered, lists[-1].start, len(lists[-1].items)), (True, 3, 2))
        self.assertEqual(_of(doc, CodeBlock)[0].code, "code line\n\nlast")

    def test_inline_formatting_and_links_survive(self):
        doc = self._roundtrip("## S\n\nplain **bold** _it_ ~~gone~~ `code` [site](https://example.com)\n")
        texts = _of(doc, Text)
        self.assertTrue(any(t.bold and t.content == "bold" for t in texts))
        self.assertTrue(any(t.italic and t.content == "it" for t in texts))
        self.assertTrue(any(t.strikethrough for t in texts))
        self.assertEqual(_of(doc, CodeInline)[0].code, "code")
        self.assertEqual(_of(doc, Link)[0].url, "https://example.com")

    def test_table_survives(self):
        doc = self._roundtrip("## S\n\n| a | b |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |\n")
        table = _of(doc, Table)[0]
        self.assertEqual([plain_text(c.children) for c in table.headers], ["a", "b"])
        self.assertEqual([[plain_text(c.children) for c in r.cells] for r in table.rows], [["1", "2"], ["3", "4"]])

    def test_a_framed_quote_returns_to_its_frame(self):
        doc = self._roundtrip("## S\n\nbody\n\n> Text box (x 10%, y 50%, w 30%, h 20%)\n>\n> on the side\n")
        quote = _of(doc, BlockQuote)[0]
        frame = parse_box_preamble(plain_text(quote.children[0].children))
        self.assertEqual((frame.x, frame.y, frame.w, frame.h), (10, 50, 30, 20))
        self.assertEqual(plain_text(quote.children[1].children), "on the side")

    def test_a_plain_quote_becomes_a_box_with_a_frame(self):
        doc = self._roundtrip("## S\n\nbody\n\n> remark\n")
        quote = _of(doc, BlockQuote)[0]
        self.assertIsNotNone(parse_box_preamble(plain_text(quote.children[0].children)))

    def test_the_first_slide_uses_the_title_layout_only_when_it_holds_little_text(self):
        path = self._save("## Opening\n\nA subtitle\n\n## List\n\n- a\n- b\n")
        names = [s.slide_layout.name for s in Presentation(path).slides]
        self.assertEqual(names, ["Title Slide", "Title and Content"])
        self.assertEqual(_headings(self._read(path)), [(2, "Opening"), (2, "List")])

    def _layouts(self, source: str) -> list[str]:
        return [s.slide_layout.name for s in Presentation(self._save(source)).slides]

    def test_a_first_level_2_slide_is_a_title_only_with_one_short_line(self):
        long_line = "word " * 20
        for body, layout in (("A subtitle", "Title Slide"), ("Ends with a period.", "Title Slide"),
                             ("One.\n\nTwo.", "Title and Content"), ("- a\n- b", "Title and Content"),
                             (long_line, "Title and Content"), ("line  \nbreak", "Title and Content")):
            with self.subTest(body):
                self.assertEqual(self._layouts(f"# Deck\n\n## Intro\n\n{body}\n\n## Next\n\nx\n")[0], layout)

    def test_a_section_with_its_own_text_keeps_the_title_layout(self):
        self.assertEqual(self._layouts("# Deck\n\nOne.\n\nTwo.\n\n## Next\n\nx\n")[0], "Title Slide")

    def test_a_bare_level_1_heading_stays_a_section_name_when_slides_follow(self):
        path = self._save("# Deck\n\n## Intro\n\nx\n")
        titles = [s.shapes.title.text_frame.text for s in Presentation(path).slides]
        self.assertEqual(titles, ["Intro"])

    def test_bare_level_1_headings_alone_get_a_title_slide_each(self):
        path = self._save("# One\n\n# Two\n")
        slides = Presentation(path).slides
        self.assertEqual([s.slide_layout.name for s in slides], ["Title Slide", "Title Slide"])
        self.assertEqual([s.shapes.title.text_frame.text for s in slides], ["One", "Two"])
        self.assertEqual([len(s.placeholders) for s in slides], [1, 1])
        self.assertEqual(_headings(self._read(path)), [(1, "One"), (1, "Two")])

    def test_a_single_bare_level_1_heading_gets_a_slide(self):
        path = self._save("# Only\n")
        self.assertEqual([s.slide_layout.name for s in Presentation(path).slides], ["Title Slide"])

    def test_sections_default_section_and_the_sections_own_slide(self):
        doc = self._roundtrip("## Before\n\n- x\n\n# A\n\nintro\n\n## In A\n\n# B\n\n## In B\n")
        path = self._save("## Before\n\n- x\n\n# A\n\nintro\n\n## In A\n\n# B\n\n## In B\n")
        sections = etree.tostring(Presentation(path).part._element, encoding="unicode")
        self.assertIn('name="Default Section"', sections)
        doc = self._read(path)
        self.assertEqual(_headings(doc), [(2, "Before"), (1, "A"), (2, "In A"), (1, "B"), (2, "In B")])
        a = next(i for i, n in enumerate(doc.children) if isinstance(n, Heading) and plain_text(n.children) == "A")
        self.assertEqual(plain_text(doc.children[a + 1].children), "intro")

    def test_a_section_with_that_name_elsewhere_is_kept(self):
        doc = self._roundtrip("# A\n\n## One\n\n# Default Section\n\n## Two\n")
        self.assertIn((1, "Default Section"), _headings(doc))

    def test_footnotes_are_speaker_notes(self):
        doc = Document(children=[
            Heading(level=2, children=[Text(content="S"), Footnote(content=[Text(content="first"), LineBreak(),
                                                                                Text(content="second")])]),
            Paragraph(children=[Text(content="body"), Footnote(content=[Text(content="aside")])]),
        ])
        path = self._render(doc)
        self.assertEqual(Presentation(path).slides[0].notes_slide.notes_text_frame.text, "first\nsecond\naside")
        self.assertEqual(len(_of(self._read(path), Footnote)), 1)

    def test_equations_are_native(self):
        doc = Document(children=[Heading(level=2, children=[Text(content="S")]),
                                 Paragraph(children=[Text(content="so "), Equation(code="x^2", inline=True)]),
                                 Equation(code=r"\frac{a}{b}", inline=False)])
        equations = _of(self._read(self._render(doc)), Equation)
        self.assertEqual([(e.code, e.inline) for e in equations], [("x^2", True), (r"\frac{a}{b}", False)])

    def test_pictures_are_embedded_and_keep_their_alt(self):
        import base64
        url = "data:image/png;base64," + base64.b64encode(_png().getvalue()).decode()
        path = self._save(f"## S\n\ntext\n\n![a red box]({url})\n")
        with zipfile.ZipFile(path) as z:
            self.assertEqual(len([n for n in z.namelist() if n.startswith("ppt/media/")]), 1)
        self.assertEqual(_of(self._read(path), Image)[0].alt, "a red box")
        path = self._save(f"## S\n\ntext\n\n![a red box]({url})\n", image_handling="discard")
        with zipfile.ZipFile(path) as z:
            self.assertFalse([n for n in z.namelist() if n.startswith("ppt/media/")])

    def test_a_missing_picture_becomes_its_alt_text(self):
        doc = self._read(self._save("## S\n\n![gone](nowhere.png)\n"))
        self.assertIn("[gone]", plain_text(_of(doc, Paragraph)[0].children))

    def test_a_crowded_slide_shrinks_its_text(self):
        body = "\n\n".join(f"Paragraph number {i} with a fair amount of words to fill the line" for i in range(30))
        path = self._save(f"## S\n\n{body}\n")
        sizes = {int(r.get("sz")) for r in Presentation(path).slides[0].element.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")}
        body = sizes - {3200}  # the title keeps its size
        self.assertTrue(body and max(body) < 2000 and min(body) >= 1000, sizes)

    def test_metadata_is_written_and_a_missing_date_is_not_invented(self):
        doc = self._roundtrip("---\ntitle: T\nauthor: A\ndate: 2026-01-02\n---\n\n## S\n")
        self.assertEqual((plain_text(doc.title), plain_text(doc.author), plain_text(doc.date)),
                         ("T", "A", "2026-01-02"))
        self.assertEqual(self._roundtrip("## S\n").date, [])

    def test_the_fixture_keeps_its_content_through_a_round_trip(self):
        from textconverter import save_to_file
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "out.pptx")
            save_to_file(str(FIXTURE), out)
            before = parse_pptx(str(FIXTURE), output_dir=tmp, image_dir_name="a")
            after = parse_pptx(out, output_dir=tmp, image_dir_name="b")
        self.assertEqual(_headings(before), _headings(after))
        for cls in (Table, Image, BlockQuote, Footnote, CodeBlock):
            self.assertEqual(len(_of(before, cls)), len(_of(after, cls)), cls.__name__)
        self.assertEqual(sum(len(b.items) for b in _of(before, ListBlock)),
                         sum(len(b.items) for b in _of(after, ListBlock)))

    def test_a_bibliography_survives(self):
        path = self._save("## S\n\nText \\cite{a}.\n\n\\bibliography{refs}\n")
        self.assertIn("\\bibliography{refs}", convert(path, "latex", textconverter.ConvertOptions(image_handling="link")))


class TestTextBoxFont(PptxOutputTest):
    """A text box tells its font size, and the writer starts from it."""

    A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def _deck(self, fill) -> str:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        body = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        body.text_frame.text = "The longest text of this slide, so that it is the body and not a box"
        box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1))
        fill(box.text_frame)
        path = str(self._out("font_box.pptx"))
        prs.save(path)
        return path

    def _font(self, path: str):
        (quote,) = _of(self._read(path), BlockQuote)
        return parse_box_preamble(plain_text(quote.children[0].children)).font

    def test_the_size_written_on_the_run(self):
        def fill(frame):
            frame.paragraphs[0].add_run().text = "text"
            frame.paragraphs[0].runs[0].font.size = Pt(24)
        self.assertEqual(self._font(self._deck(fill)), 24)

    def test_the_size_used_by_most_characters(self):
        def fill(frame):
            for text, size in (("a long stretch of text", 12), ("big", 40)):
                run = frame.paragraphs[0].add_run()
                run.text = text
                run.font.size = Pt(size)
        self.assertEqual(self._font(self._deck(fill)), 12)

    def test_the_size_from_the_list_style_of_the_shape(self):
        def fill(frame):
            frame.text = "text"
            style = frame._txBody.find(f"{self.A_NS}lstStyle")
            level = etree.SubElement(style, f"{self.A_NS}lvl1pPr")
            etree.SubElement(level, f"{self.A_NS}defRPr").set("sz", "1600")
        self.assertEqual(self._font(self._deck(fill)), 16)

    def test_the_size_from_the_defaults_of_the_presentation(self):
        self.assertEqual(self._font(self._deck(lambda frame: setattr(frame, "text", "text"))), 18)

    def test_the_writer_uses_it(self):
        doc = self._save("## S\n\nbody\n\n> Text box (x 10%, y 50%, w 80%, h 40%, font 28pt)\n>\n> short\n")
        sizes = [int(r.get("sz")) for shape in Presentation(doc).slides[0].shapes if shape.has_text_frame
                 and "short" in shape.text_frame.text for r in shape._element.iter(f"{self.A_NS}rPr")]
        self.assertEqual(set(sizes), {2800})

    def test_the_writer_still_shrinks_what_does_not_fit(self):
        text = " ".join(["word"] * 200)
        doc = self._save(f"## S\n\nbody\n\n> Text box (x 10%, y 50%, w 20%, h 10%, font 28pt)\n>\n> {text}\n")
        sizes = [int(r.get("sz")) for shape in Presentation(doc).slides[0].shapes if shape.has_text_frame
                 and "word" in shape.text_frame.text for r in shape._element.iter(f"{self.A_NS}rPr")]
        self.assertLess(max(sizes), 2800)

    def test_the_font_comes_back(self):
        doc = self._roundtrip("## S\n\nbody\n\n> Text box (x 10%, y 50%, w 80%, h 40%, font 28pt)\n>\n> short\n")
        self.assertEqual(parse_box_preamble(plain_text(_of(doc, BlockQuote)[0].children[0].children)).font, 28)

    def test_a_box_without_a_place_on_the_slide_flows_under_the_content(self):
        for preamble in ("Text box (x 10%, w 30%, h 20%, font 11pt)", "Text box (C4:M13, font 11pt)"):
            with self.subTest(preamble):
                path = self._save(f"## S\n\nbody\n\n> {preamble}\n>\n> side note\n")
                texts = [s.text_frame.text for s in Presentation(path).slides[0].shapes if s.has_text_frame]
                self.assertIn("side note", texts)
                self.assertFalse(any("Text box" in t for t in texts))


class TestPptxOutputThroughTheApi(BaseConverterTest):

    def test_text_conversion_cannot_produce_pptx(self):
        with self.assertRaisesRegex(ValueError, "to_format"):
            textconverter.convert_text("# a", "markdown", "pptx")
        with self.assertRaisesRegex(ValueError, "to_format"):
            convert(str(FIXTURE), "pptx")

    def test_output_format_is_inferred(self):
        from textconverter.formats import OUTPUT_FORMATS, infer_output_format
        self.assertEqual(infer_output_format("a.pptx"), "pptx")
        self.assertEqual(OUTPUT_FORMATS["pptx"].auto_images, "embed")

    def test_a_failed_conversion_leaves_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = os.path.join(tmp, "bad.pptx")
            with open(bad, "w", encoding="utf-8") as f:
                f.write("no")
            with self.assertRaises(ValueError):
                textconverter.save_to_file(bad, os.path.join(tmp, "out.pptx"))
            self.assertEqual(os.listdir(tmp), ["bad.pptx"])

    def test_the_source_is_never_overwritten(self):
        with self.assertRaisesRegex(ValueError, "source file itself"):
            textconverter.save_to_file(str(FIXTURE), str(FIXTURE))
