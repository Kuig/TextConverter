"""Word (.docx) input: the real fixture Test_docx.docx and documents built with python-docx.

Usage: python -m unittest discover -s tests -p "test_docx.py"
"""
from __future__ import annotations
import io
import os
import tempfile
import zipfile
from unittest.mock import patch

import docx
from lxml import etree

import textconverter
from textconverter.ast import (
    BlockQuote, CodeBlock, CodeInline, Equation, Footnote, Heading, Image, LineBreak, Link, ListBlock,
    Paragraph, Table, Text, walk,
)
from textconverter.parsers._omml import M_NS, omml_to_latex
from textconverter.parsers.docx_parser import parse_docx
from textconverter.renderers._common import plain_text
from _support import BaseConverterTest, setUpModule, tearDownModule, INPUT_DIR

FIXTURE = INPUT_DIR / "Test_docx.docx"


def _of(doc, cls):
    return [n for n in walk(doc) if isinstance(n, cls)]


def _png() -> io.BytesIO:
    from PIL import Image as PilImage
    buffer = io.BytesIO()
    PilImage.new("RGB", (4, 3), "red").save(buffer, "PNG")
    buffer.seek(0)
    return buffer


class DocxFixtureTest(BaseConverterTest):
    """Parses Test_docx.docx once; the tests read the same AST."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.tmp.cleanup)
        with patch("textconverter.parsers.docx_parser.log_warning") as warn:
            cls.doc = parse_docx(str(FIXTURE), output_dir=cls.tmp.name, image_dir_name="img")
        cls.warnings = [call.args[0] for call in warn.call_args_list]


class TestCodeDetection(BaseConverterTest):
    CODE = ["import os", "def f(x):", "    y = x + 1", "    if y:", "        return y", "    else:", "        return 0"]
    PROSE = "Questa è una frase normale, con una virgola e (tra parentesi) un inciso."

    def _save(self, tmp: str, *lines: str) -> str:
        document = docx.Document()
        for line in lines:
            document.add_paragraph(line)
        path = os.path.join(tmp, "code.docx")
        document.save(path)
        return path

    def test_plain_paragraphs_of_code_become_one_block_with_the_flag(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._save(tmp, self.PROSE, *self.CODE, self.PROSE)
            markdown = textconverter.convert(path, "markdown", textconverter.ConvertOptions(code_parsing=True, image_handling="discard"))
        self.assertIn("```python\n" + "\n".join(self.CODE) + "\n```", markdown)
        self.assertEqual(markdown.count(self.PROSE), 2)

    def test_nothing_changes_without_the_flag(self):
        with tempfile.TemporaryDirectory() as tmp:
            markdown = textconverter.convert(self._save(tmp, *self.CODE), "markdown", textconverter.ConvertOptions(image_handling="discard"))
        self.assertNotIn("```", markdown)

    def test_prose_stays_prose(self):
        lines = [self.PROSE, "Vedi anche: la sezione seguente (pag. 3).", "Esempio:", "Il costo è 3 = 2 + 1."]
        with tempfile.TemporaryDirectory() as tmp:
            markdown = textconverter.convert(self._save(tmp, *lines), "markdown",
                               textconverter.ConvertOptions(code_parsing=True, image_handling="discard"))
        self.assertNotIn("```", markdown)

    def test_code_in_a_monospaced_font_stays_one_block(self):
        document = docx.Document()
        for line in ("x = 1", "y = 2"):
            document.add_paragraph().add_run(line).font.name = "Consolas"
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "mono.docx")
            document.save(path)
            options = textconverter.ConvertOptions(image_handling="discard")
            self.assertEqual(textconverter.convert(path, "markdown", options),
                             textconverter.convert(path, "markdown", textconverter.ConvertOptions(code_parsing=True, image_handling="discard")))


class TestFixtureStructure(DocxFixtureTest):

    def test_title_and_metadata(self):
        """The Title paragraph wins over the file properties; author and date come from the properties."""
        self.assertEqual(plain_text(self.doc.title), "Documento di test")
        self.assertEqual(plain_text(self.doc.author), "Giorgio Presti")
        self.assertEqual(plain_text(self.doc.date), "2026-10-30")

    def test_subtitle_is_a_paragraph_and_headings_have_levels(self):
        first = self.doc.children[0]
        self.assertIsInstance(first, Paragraph)
        self.assertEqual(plain_text(first.children), "Questo è un sottotitolo")
        levels = [(h.level, plain_text(h.children)) for h in _of(self.doc, Heading)]
        self.assertIn((1, "Testo:"), levels)
        self.assertIn((2, "Paragrafi:"), levels)
        self.assertTrue(all(text.strip() for _, text in levels), "empty headings must be dropped")

    def test_inline_formatting_links_and_scripts(self):
        texts = _of(self.doc, Text)
        self.assertTrue(any(t.strikethrough for t in texts))
        self.assertTrue(any(t.bold and t.content.startswith("Lorem") for t in texts))
        self.assertEqual({t.script for t in texts if t.script}, {"sup", "sub"})
        link = _of(self.doc, Link)[0]
        self.assertEqual(link.url, "https://www.google.com")
        self.assertEqual(plain_text(link.content), "link")

    def test_tabs_become_spaces(self):
        self.assertFalse(any("\t" in t.content for t in _of(self.doc, Text)))

    def test_quote_and_text_box(self):
        firsts = [plain_text(q.children[0].children) if isinstance(q.children[0], Paragraph) else ""
                  for q in _of(self.doc, BlockQuote)]
        self.assertIn("Nuovo paragrafo come stile citazione", firsts)
        boxes = [q for q in _of(self.doc, BlockQuote) if len(q.children) > 1
                 and plain_text(q.children[1].children).startswith("Casella di testo")]
        self.assertEqual(len(boxes), 1, "the text box is read once, not once per AlternateContent branch")

    def test_a_text_box_gives_its_frame_and_font(self):
        (box,) = [q for q in _of(self.doc, BlockQuote) if len(q.children) > 1
                  and plain_text(q.children[1].children).startswith("Casella di testo")]
        # The fixture anchors it to a paragraph, so its height on the page is not known
        self.assertEqual(plain_text(box.children[0].children), "Text box (x 44%, w 23%, h 4%, font 11pt)")

    def test_code_block_and_inline_code(self):
        blocks = _of(self.doc, CodeBlock)
        self.assertEqual(len(blocks), 1)
        self.assertEqual(blocks[0].code, "for (int i = 0; i < k; ++i)\n{\n\tobject.print(i);\n}")
        self.assertIn("inline", [c.code for c in _of(self.doc, CodeInline)])

    def test_formula_is_latex(self):
        equation = _of(self.doc, Equation)[0]
        self.assertFalse(equation.inline)
        self.assertIn(r"\sum_{n=1}^\infty", equation.code)
        self.assertIn(r"\frac{n\pi x}{L}", equation.code)

    def test_footnote(self):
        notes = _of(self.doc, Footnote)
        self.assertEqual(len(notes), 1)
        self.assertEqual(plain_text(notes[0].content), "Ecco la nota, formattata")

    def test_image_figure_and_caption(self):
        images = _of(self.doc, Image)
        self.assertEqual(len(images), 1)
        self.assertEqual(images[0].src, "img/image_1.jpeg")
        self.assertEqual(plain_text(images[0].caption), "Esempio di immagine con didascalia")
        self.assertTrue(os.path.isfile(os.path.join(self.tmp.name, "img", "image_1.jpeg")))

    def test_warning_lists_what_was_dropped(self):
        self.assertEqual(len(self.warnings), 1)
        for what in ("underlined text", "highlighted text", "shape", "header", "footer"):
            self.assertIn(what, self.warnings[0])


class TestFixtureListsAndTables(DocxFixtureTest):

    def test_mixed_nested_lists(self):
        lists = [b for b in self.doc.children if isinstance(b, ListBlock)]
        self.assertEqual([lst.ordered for lst in lists], [False, True, True])
        bullets, numbered, lettered = lists
        self.assertEqual([plain_text(i.children[0].children) for i in bullets.items][:3], ["A", "B e c e d", "EFG"])
        nested = bullets.items[2].children[-1]
        self.assertIsInstance(nested, ListBlock)
        self.assertEqual(len(nested.items), 3)
        self.assertFalse(numbered.items[1].children[-1].ordered, "bullets inside a numbered list")
        self.assertTrue(lettered.items[0].children[-1].ordered)

    def test_tables_headers_and_merged_cells(self):
        first, second = _of(self.doc, Table)
        self.assertEqual([plain_text(c.children) for c in first.headers], ["Prova 11", "12"])
        self.assertEqual(len(first.rows), 1)
        self.assertEqual([plain_text(c.children) for c in second.headers], ["celle unite", "", "Non unita"])
        self.assertEqual([plain_text(c.children) for c in second.rows[1].cells], ["DDati", "Altre celle unite", ""])
        self.assertTrue(any(isinstance(n, LineBreak) for n in walk(first.rows[0].cells[1])))

    def test_captions_lose_label_and_number(self):
        first, second = _of(self.doc, Table)
        self.assertEqual(first.caption, [])
        self.assertEqual(plain_text(second.caption),
                         "In realtà è la due perchè la 1 non ha didascalia ma va bene")


class TestGeneratedDocuments(BaseConverterTest):
    """Small documents built with python-docx."""

    def _parse(self, build, **kwargs):
        document = docx.Document()
        build(document)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "doc.docx")
            document.save(path)
            return parse_docx(path, output_dir=tmp, image_dir_name="img", **kwargs)

    def test_style_based_lists_nest_by_style_name(self):
        def build(d):
            d.add_paragraph("a", style="List Bullet")
            d.add_paragraph("b", style="List Bullet 2")
            d.add_paragraph("c", style="List Bullet")
            d.add_paragraph("one", style="List Number")
        doc = self._parse(build)
        bullets, numbered = doc.children
        self.assertEqual(len(bullets.items), 2)
        self.assertEqual(len(bullets.items[0].children), 2)
        self.assertTrue(numbered.ordered)

    def test_task_items_and_reference_macros(self):
        def build(d):
            d.add_paragraph("☑ done", style="List Bullet")
            d.add_paragraph("☐ todo", style="List Bullet")
            d.add_paragraph("As in \\cite{knuth} and \\ref{fig:a}.")
        doc = self._parse(build)
        self.assertEqual([i.checked for i in doc.children[0].items], [True, False])
        kinds = {type(n).__name__ for n in walk(doc.children[1])}
        self.assertTrue({"Citation", "Reference"} <= kinds)

    def test_caption_before_a_table_and_after_a_figure(self):
        def build(d):
            caption = d.add_paragraph("Totals", style="Caption")
            caption.add_run().add_text("")
            table = d.add_table(rows=1, cols=1)
            table.rows[0].cells[0].text = "x"
            d.add_paragraph().add_run().add_picture(_png())
            d.add_paragraph("A red square", style="Caption")
        doc = self._parse(build)
        table = _of(doc, Table)[0]
        self.assertEqual(plain_text(table.caption), "Totals")
        self.assertEqual(plain_text(_of(doc, Image)[0].caption), "A red square")

    def test_display_equation_style_and_environment(self):
        def build(d):
            d.styles.add_style("Equation Display", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
            d.add_paragraph("\\begin{align}a&=b\\end{align}", style="Equation Display")
        equation = self._parse(build).children[0]
        self.assertEqual((equation.code, equation.env, equation.inline), ("a&=b", "align", False))

    def test_images_can_be_left_out_and_get_generated_names(self):
        def build(d):
            d.add_paragraph().add_run().add_picture(_png())
        doc = self._parse(build)
        self.assertEqual(_of(doc, Image)[0].src, "img/image_1.png")
        self.assertEqual(_of(self._parse(build, write_images=False), Image), [])

    def test_same_image_twice_is_one_file(self):
        def build(d):
            for _ in range(2):
                d.add_paragraph().add_run().add_picture(_png())
        doc = self._parse(build)
        self.assertEqual({i.src for i in _of(doc, Image)}, {"img/image_1.png"})

    def test_not_a_word_document_is_a_value_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = os.path.join(tmp, "bad.docx")
            with open(bad, "w", encoding="utf-8") as f:
                f.write("not a zip")
            with self.assertRaisesRegex(ValueError, "Not a readable Word"):
                parse_docx(bad)
            empty = os.path.join(tmp, "empty.docx")
            with zipfile.ZipFile(empty, "w") as z:
                z.writestr("hello.txt", "x")
            with self.assertRaisesRegex(ValueError, "Not a readable Word"):
                parse_docx(empty)


class TestOmml(BaseConverterTest):

    def _latex(self, body: str) -> str:
        xml = f'<m:oMath xmlns:m="{M_NS}">{body}</m:oMath>'
        return omml_to_latex(etree.fromstring(xml))

    @staticmethod
    def _r(text: str) -> str:
        return f"<m:r><m:t>{text}</m:t></m:r>"

    def test_scripts_fractions_and_radicals(self):
        r = self._r
        self.assertEqual(self._latex(f"<m:sSup><m:e>{r('x')}</m:e><m:sup>{r('2')}</m:sup></m:sSup>"), "x^2")
        self.assertEqual(self._latex(f"<m:f><m:num>{r('a')}</m:num><m:den>{r('b')}</m:den></m:f>"), r"\frac{a}{b}")
        self.assertEqual(self._latex(f"<m:rad><m:radPr><m:degHide m:val='1'/></m:radPr><m:deg/><m:e>{r('x')}</m:e></m:rad>"),
                         r"\sqrt{x}")
        self.assertEqual(self._latex(f"<m:rad><m:deg>{r('3')}</m:deg><m:e>{r('x')}</m:e></m:rad>"), r"\sqrt[3]{x}")

    def test_greek_letters_are_kept_apart_from_following_letters(self):
        self.assertEqual(self._latex(self._r("αx")), r"\alpha x")

    def test_matrix_and_delimiters(self):
        r = self._r
        row = f"<m:mr><m:e>{r('a')}</m:e><m:e>{r('b')}</m:e></m:mr>"
        self.assertEqual(self._latex(f"<m:m>{row}{row}</m:m>"),
                         r"\begin{matrix} a & b \\ a & b \end{matrix}")
        self.assertEqual(self._latex(f"<m:d><m:dPr><m:begChr m:val='['/><m:endChr m:val=']'/></m:dPr><m:e>{r('x')}</m:e></m:d>"),
                         r"\left[x\right]")

    def test_unknown_elements_contribute_their_text(self):
        self.assertEqual(self._latex(f"<m:phant><m:e>{self._r('x')}</m:e></m:phant>"), "x")


class TestDocxThroughTheApi(BaseConverterTest):

    def test_convert_to_markdown_and_html(self):
        with tempfile.TemporaryDirectory() as tmp:
            markdown = textconverter.convert(str(FIXTURE), "markdown", output_dir=tmp, image_dir_name="images")
            self.assertIn("# Elenchi:", markdown)
            self.assertIn("![Foto](images/image_1.jpeg)", markdown)
            self.assertTrue(os.path.isfile(os.path.join(tmp, "images", "image_1.jpeg")))
            html = textconverter.convert(str(FIXTURE), "html")
            self.assertIn("data:image/jpeg;base64,", html)

    def test_text_conversion_needs_a_file(self):
        with self.assertRaisesRegex(ValueError, "needs a file"):
            textconverter.convert_text("x", "docx", "markdown")

    def test_format_is_inferred_and_documented(self):
        from textconverter.formats import infer_input_format, infer_input_format_from_media_type
        self.assertEqual(infer_input_format("a.DOCX"), "docx")
        self.assertEqual(infer_input_format_from_media_type(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), "docx")


class DocxOutputTest(BaseConverterTest):
    """Helpers for tests that write a Word file."""

    def _save(self, source: str, **options) -> str:
        """Convert Markdown ``source`` to a .docx in a temporary folder; return the .docx path."""
        from textconverter import ConvertOptions, save_to_file
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        src = os.path.join(tmp.name, "in.md")
        with open(src, "w", encoding="utf-8") as f:
            f.write(source)
        out = os.path.join(tmp.name, "out.docx")
        save_to_file(src, out, ConvertOptions(**options))
        self.assertEqual(sorted(os.listdir(tmp.name)), ["in.md", "out.docx"], "no image folder or temp file left")
        return out

    def _roundtrip(self, source: str):
        out = self._save(source)
        with tempfile.TemporaryDirectory() as tmp:
            return parse_docx(out, output_dir=tmp, image_dir_name="img")


def _data_url(blob: bytes, mime: str = "image/png") -> str:
    import base64
    return f"data:{mime};base64," + base64.b64encode(blob).decode()


class TestDocxRenderer(DocxOutputTest):

    def test_fixture_survives_a_docx_round_trip(self):
        from textconverter import save_to_file
        from textconverter.renderers.json_renderer import render_json
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "out.docx")
            save_to_file(str(FIXTURE), out)
            again = parse_docx(out, output_dir=tmp, image_dir_name="img")
            first = parse_docx(str(FIXTURE), output_dir=tmp, image_dir_name="img")
            normalized = lambda doc: render_json(doc).replace("image_1.jpg", "image_1.jpeg")  # noqa: E731
            self.assertEqual(normalized(again), normalized(first))

    def test_text_formatting_links_and_metadata(self):
        doc = self._roundtrip(
            "---\ntitle: My title\nauthor: Ann\ndate: 2026-01-02\n---\n\n"
            "# Head\n\nSome **bold**, _italic_, ~~gone~~, `code`, x<sup>2</sup>, H<sub>2</sub>O "
            "and [a link](https://example.com/a).\n\nA second\\\nline.")
        self.assertEqual((plain_text(doc.title), plain_text(doc.author), plain_text(doc.date)),
                         ("My title", "Ann", "2026-01-02"))
        texts = {t.content: t for t in _of(doc, Text)}
        self.assertTrue(texts["bold"].bold and texts["italic"].italic and texts["gone"].strikethrough)
        self.assertEqual([t.script for t in _of(doc, Text) if t.content == "2"], ["sup", "sub"])
        self.assertEqual([c.code for c in _of(doc, CodeInline)], ["code"])
        self.assertEqual(_of(doc, Link)[0].url, "https://example.com/a")
        self.assertTrue(any(isinstance(n, LineBreak) for n in walk(doc)))

    def test_without_a_date_the_templates_date_is_not_written(self):
        out = self._save("# Plain")
        self.assertIsNone(docx.Document(out).core_properties.created)
        self.assertEqual(plain_text(parse_docx(out, image_dir_name="x").date), "")

    def test_lists_keep_kind_start_nesting_and_task_state(self):
        doc = self._roundtrip(
            "1. one\n2. two\n   - nested\n   - bullets\n\ntext\n\n3. three\n\n- [x] done\n- [ ] todo\n")
        first, second, tasks = [b for b in doc.children if isinstance(b, ListBlock)]
        self.assertEqual((first.ordered, first.start, len(first.items)), (True, 1, 2))
        self.assertFalse(first.items[1].children[-1].ordered)
        self.assertEqual((second.start, plain_text(second.items[0].children[0].children)), (3, "three"))
        self.assertEqual([i.checked for i in tasks.items], [True, False])

    def test_quote_alert_code_and_rule(self):
        doc = self._roundtrip("> quoted\n\nbetween\n\n> [!NOTE]\n> careful\n\n```python\nx = 1\n\ty = 2\n```\n\n---\n\nend")
        quotes = _of(doc, BlockQuote)
        self.assertEqual(plain_text(quotes[0].children[0].children), "quoted")
        self.assertIn("Note", plain_text(quotes[1].children[0].children))
        self.assertEqual(_of(doc, CodeBlock)[0].code, "x = 1\n\ty = 2")
        self.assertTrue(any(type(n).__name__ == "HorizontalRule" for n in doc.children))

    def test_abstract(self):
        from textconverter.ast import Abstract
        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, "a.tex")
            with open(src, "w", encoding="utf-8") as f:
                f.write("\\begin{document}\\begin{abstract}Short text.\\end{abstract}\\end{document}")
            out = os.path.join(tmp, "a.docx")
            textconverter.save_to_file(src, out)
            abstract = _of(parse_docx(out, output_dir=tmp), Abstract)[0]
        self.assertEqual(plain_text(abstract.children).strip(), "Short text.")

    def test_tables_and_captions_with_labels(self):
        doc = self._roundtrip("Table: Totals \\label{tab:t}\n\n| A | B |\n|---|---|\n| **x** | y |\n\n"
                              f"![pic]({_data_url(_png().getvalue())})\n\nFigure: A pic")
        table = _of(doc, Table)[0]
        self.assertEqual([plain_text(c.children) for c in table.headers], ["A", "B"])
        self.assertTrue(table.rows[0].cells[0].children[0].bold)
        self.assertEqual((plain_text(table.caption), table.label), ("Totals", "tab:t"))
        image = _of(doc, Image)[0]
        self.assertEqual((plain_text(image.caption), image.alt), ("A pic", "pic"))

    def test_captions_are_numbered_by_word(self):
        out = self._save("Table: One\n\n| A |\n|---|\n| 1 |\n\nTable: Two\n\n| A |\n|---|\n| 2 |\n")
        with zipfile.ZipFile(out) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        self.assertEqual(xml.count("SEQ Table"), 2)

    def test_equations_are_native_word_equations(self):
        source = "Inline $a_1+\\frac{b}{c}$ here.\n\n$$\\begin{align}x&=1\\end{align}$$\n"
        out = self._save(source)
        with zipfile.ZipFile(out) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        self.assertIn("<m:oMathPara", xml)
        self.assertEqual(xml.count("<m:oMath>") + xml.count("<m:oMath "), 2)
        self.assertNotIn("a_1", xml, "the formula is structure, not LaTeX text")
        doc = self._roundtrip(source)
        inline, display = _of(doc, Equation)
        self.assertEqual((inline.code, inline.inline), ("a_1+\\frac{b}{c}", True))
        self.assertEqual((display.code, display.env, display.inline), ("x&=1", "align", False))

    def test_a_formula_that_cannot_be_translated_stays_latex_text(self):
        with patch("textconverter.renderers.docx_renderer.log_warning") as warn:
            doc = self._roundtrip("Odd $\\unknowncmd{x}$ and\n\n$$\\unknowncmd{y}$$\n")
        self.assertEqual(warn.call_count, 2)
        inline, display = _of(doc, Equation)
        self.assertEqual((inline.code, inline.inline), ("\\unknowncmd{x}", True))
        self.assertEqual((display.code, display.inline), ("\\unknowncmd{y}", False))

    def test_footnotes_are_native_and_references_stay_macros(self):
        out = self._save("A claim[^1] and \\cite{key}, see \\ref{sec:a}.\n\n[^1]: The **note** text.")
        with zipfile.ZipFile(out) as z:
            self.assertIn("word/footnotes.xml", z.namelist())
            self.assertIn("footnoteReference", z.read("word/document.xml").decode("utf-8"))
        doc = parse_docx(out, image_dir_name="x")
        self.assertEqual(plain_text(_of(doc, Footnote)[0].content), "The note text.")
        self.assertTrue({"Citation", "Reference"} <= {type(n).__name__ for n in walk(doc)})

    def test_images_are_embedded_and_unreadable_ones_become_text(self):
        out = self._save(f"![ok]({_data_url(_png().getvalue())})\n\n![gone](missing.png)\n")
        doc = parse_docx(out, output_dir=os.path.dirname(out), image_dir_name="img")
        self.assertEqual([i.alt for i in _of(doc, Image)], ["ok"])
        self.assertIn("[gone]", plain_text(doc.children))

    def test_webp_is_converted_to_png(self):
        from PIL import Image as PilImage
        buffer = io.BytesIO()
        PilImage.new("RGB", (8, 8), "blue").save(buffer, "WEBP")
        doc = self._roundtrip(f"![w]({_data_url(buffer.getvalue(), 'image/webp')})")
        self.assertEqual(len(_of(doc, Image)), 1)

    def test_images_are_dropped_with_discard(self):
        out = self._save(f"![cap]({_data_url(_png().getvalue())})", image_handling="discard")
        with zipfile.ZipFile(out) as z:
            self.assertFalse([n for n in z.namelist() if n.startswith("word/media")])

    def test_a_jpeg_without_a_resolution_is_embedded(self):
        """JFIF allows 0 dpi; python-docx then divides by zero when it sizes the picture."""
        from PIL import Image as PilImage
        buffer = io.BytesIO()
        PilImage.new("RGB", (8, 8), "blue").save(buffer, "JPEG")
        blob = bytearray(buffer.getvalue())
        blob[13:18] = b"\x01\x00\x00\x00\x00"  # units: dpi, density: 0 by 0
        blob = bytes(blob)
        from docx.image.image import Image as DocxImage
        self.assertEqual(DocxImage.from_blob(blob).horz_dpi, 0)
        with patch("textconverter.renderers.docx_renderer.log_warning") as warn:
            doc = self._roundtrip(f"![j]({_data_url(blob, 'image/jpeg')})")
        warn.assert_not_called()
        self.assertEqual(len(_of(doc, Image)), 1)

    def test_code_language_survives(self):
        for language in ("python", "excel"):
            with self.subTest(language):
                (block,) = _of(self._roundtrip(f"```{language}\nx = 1\ny = 2\n```\n"), CodeBlock)
                self.assertEqual((block.code, block.language), ("x = 1\ny = 2", language))

    def test_code_without_language_has_no_control(self):
        with zipfile.ZipFile(self._save("```\nx = 1\n```\n")) as z:
            self.assertNotIn(b"<w:sdt>", z.read("word/document.xml"))

    def test_a_bibliography_survives(self):
        out = self._save("Text \\cite{a}.\n\n\\bibliographystyle{plain}\n\n\\bibliography{refs}\n\n\\bibitem{a} One.\n")
        doc = textconverter.convert(out, "json", textconverter.ConvertOptions(image_handling="link"))
        self.assertIn('"Bibliography"', doc)
        self.assertIn("\\bibliography{refs}", textconverter.convert(out, "latex", textconverter.ConvertOptions(image_handling="link")))


class TestDocxTextBoxes(DocxOutputTest):
    """A quote headed by a frame is a floating text box; reading it gives the same quote."""

    SOURCE = "Before\n\n> Text box (x 10%, y 50%, w 30%, h 20%, font 14pt)\n>\n> **Side** text\n>\n> Second\n\nAfter\n"

    def _xml(self, path: str) -> str:
        with zipfile.ZipFile(path) as z:
            return z.read("word/document.xml").decode("utf-8")

    def test_it_is_a_real_text_box(self):
        xml = self._xml(self._save(self.SOURCE))
        self.assertIn("<w:txbxContent", xml)
        self.assertIn('relativeFrom="page"', xml)
        self.assertIn('txBox="1"', xml)
        self.assertNotIn("Text box (", xml)
        self.assertRegex(xml, r'<wp:docPr id="\d+"')

    def test_the_font_size_is_applied(self):
        self.assertIn('<w:sz w:val="28"/>', self._xml(self._save(self.SOURCE)))

    def test_the_round_trip_is_identical(self):
        from textconverter.parsers._office import parse_box_preamble
        doc = self._roundtrip(self.SOURCE)
        (quote,) = _of(doc, BlockQuote)
        frame = parse_box_preamble(plain_text(quote.children[0].children))
        self.assertEqual((frame.x, frame.y, frame.w, frame.h, frame.font), (10, 50, 30, 20, 14))
        self.assertEqual(plain_text(quote.children[1].children), "Side text")
        self.assertEqual(plain_text(quote.children[2].children), "Second")
        self.assertEqual([type(n) for n in doc.children], [Paragraph, BlockQuote, Paragraph])

    def test_a_box_without_a_page_height_stays_at_its_paragraph(self):
        xml = self._xml(self._save("Before\n\n> Text box (x 10%, w 30%, h 20%, font 11pt)\n>\n> text\n"))
        self.assertIn('<wp:positionV relativeFrom="paragraph">', xml)
        self.assertEqual(plain_text(_of(self._roundtrip("Before\n\n> Text box (x 10%, w 30%, h 20%, font 11pt)\n>\n> text\n"),
                                        BlockQuote)[0].children[0].children),
                         "Text box (x 10%, w 30%, h 20%, font 11pt)")

    def test_a_box_at_the_start_of_the_document(self):
        doc = self._roundtrip("> Text box (x 10%, y 50%, w 30%, h 20%)\n>\n> first\n")
        self.assertEqual(len(_of(doc, BlockQuote)), 1)

    def test_a_box_of_cells_stays_a_quote(self):
        xml = self._xml(self._save("> Text box (C4:M13, font 11pt)\n>\n> from a sheet\n"))
        self.assertNotIn("<w:txbxContent", xml)
        self.assertIn("from a sheet", xml)

    def test_an_ordinary_quote_stays_a_quote(self):
        self.assertNotIn("<w:txbxContent", self._xml(self._save("> quoted\n")))

    def test_lists_and_tables_go_inside_the_box(self):
        doc = self._roundtrip("> Text box (x 10%, y 50%, w 30%, h 20%)\n>\n> - a\n> - b\n>\n> | x |\n> |---|\n> | 1 |\n")
        (quote,) = _of(doc, BlockQuote)
        self.assertEqual(len(_of(quote, ListBlock)), 1)
        self.assertEqual(len(_of(quote, Table)), 1)
        self.assertEqual(len(_of(doc, Table)), 1)

    def test_the_font_comes_from_the_style_when_runs_do_not_tell(self):
        from textconverter.parsers._office import parse_box_preamble
        with tempfile.TemporaryDirectory() as tmp:
            first = parse_docx(str(FIXTURE), output_dir=tmp, image_dir_name="img")
        box = next(q for q in _of(first, BlockQuote) if len(q.children) > 1
                   and plain_text(q.children[1].children).startswith("Casella"))
        self.assertEqual(parse_box_preamble(plain_text(box.children[0].children)).font, 11)

    def test_the_dominant_size_wins(self):
        document = docx.Document()
        document.add_paragraph("a")
        path = None
        with tempfile.TemporaryDirectory() as tmp:
            source = os.path.join(tmp, "in.md")
            with open(source, "w", encoding="utf-8") as f:
                f.write(self.SOURCE)
            path = os.path.join(tmp, "box.docx")
            textconverter.save_to_file(source, path)
            edited = docx.Document(path)
            run = edited.element.body.xpath(".//w:txbxContent//w:r")[0]
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            props = run.get_or_add_rPr()
            for old in props.findall(qn("w:sz")):
                props.remove(old)
            size = OxmlElement("w:sz")
            size.set(qn("w:val"), "72")
            props.append(size)
            edited.save(path)
            doc = parse_docx(path, output_dir=tmp, image_dir_name="img")
        from textconverter.parsers._office import parse_box_preamble
        self.assertEqual(parse_box_preamble(plain_text(_of(doc, BlockQuote)[0].children[0].children)).font, 14)


class TestDocxProperties(DocxOutputTest):

    def _props(self, source: str):
        return docx.Document(self._save(source)).core_properties

    def test_long_metadata_is_cut_at_a_word(self):
        author = " ".join(f"Author{i}" for i in range(60))
        with patch("textconverter.renderers._office.log_warning") as warn:
            props = self._props(f'---\ntitle: "{author}"\nauthor: "{author}"\n---\n\ntext\n')
        self.assertLessEqual(len(props.author), 255)
        self.assertLessEqual(len(props.title), 255)
        self.assertTrue(author.startswith(props.author + " "))
        warn.assert_called()

    def test_the_fixture_that_crashed_is_converted(self):
        from textconverter import ConvertOptions, save_to_file
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "weird.docx")
            save_to_file(str(INPUT_DIR / "Weird Journal TEX.tex"), out, ConvertOptions(image_handling="discard"))
            self.assertLessEqual(len(docx.Document(out).core_properties.author), 255)

    def test_lines_formulas_and_notes(self):
        tex = ("\\title{A \\thanks{Funded} title}\n\\author{Ann$^{1}$\\\\Dept}\n"
               "\\begin{document}\\maketitle\ntext\n\\end{document}\n")
        from textconverter import ConvertOptions, save_to_file
        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, "a.tex")
            with open(src, "w", encoding="utf-8") as f:
                f.write(tex)
            out = os.path.join(tmp, "a.docx")
            save_to_file(src, out, ConvertOptions(image_handling="discard"))
            props = docx.Document(out).core_properties
            self.assertEqual((props.title, props.author), ("A  title".replace("  ", " "), "Ann$^{1}$; Dept"))


class TestDiscardedCaptions(BaseConverterTest):

    def test_a_caption_goes_with_its_image(self):
        document = docx.Document()
        document.add_picture(_png())
        document.add_paragraph("Figura 1: La didascalia", style="Caption")
        document.add_paragraph("Dopo")
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "c.docx")
            document.save(path)
            md = textconverter.convert(path, "markdown", textconverter.ConvertOptions(image_handling="discard"))
        self.assertNotIn("didascalia", md)
        self.assertIn("Dopo", md)


class TestDocxOutputThroughTheApi(BaseConverterTest):

    def test_text_conversion_cannot_produce_docx(self):
        with self.assertRaisesRegex(ValueError, "to_format"):
            textconverter.convert_text("# a", "markdown", "docx")
        with self.assertRaisesRegex(ValueError, "to_format"):
            textconverter.convert(str(FIXTURE), "docx")

    def test_output_format_is_inferred(self):
        from textconverter.formats import infer_output_format, OUTPUT_FORMATS
        self.assertEqual(infer_output_format("a.docx"), "docx")
        self.assertEqual(OUTPUT_FORMATS["docx"].auto_images, "embed")

    def test_a_failed_conversion_leaves_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = os.path.join(tmp, "bad.docx")
            with open(bad, "w", encoding="utf-8") as f:
                f.write("no")
            with self.assertRaises(ValueError):
                textconverter.save_to_file(bad, os.path.join(tmp, "out.docx"))
            self.assertEqual(os.listdir(tmp), ["bad.docx"])
            with self.assertRaisesRegex(ValueError, "from_format"):
                textconverter.save_to_file(bad, os.path.join(tmp, "o.docx"), from_format="nope")

    def test_the_source_is_never_overwritten(self):
        with self.assertRaisesRegex(ValueError, "source file itself"):
            textconverter.save_to_file(str(FIXTURE), str(FIXTURE))


class TestLatexToOmml(BaseConverterTest):
    """Formulas written as OMML read back as the same LaTeX."""

    FORMULAS = (
        r"a_1+\frac{b}{c}",
        r"x^2+y_i^2",
        r"\sqrt{x}+\sqrt[3]{y}",
        r"\sum_{n=1}^\infty a_n",
        r"\int_0^1 f\left(x\right)",
        r"\left[x+1\right]",
        r"\alpha x+\beta",
        r"\hat{x}+\overline{y}+\vec{v}",
        r"\cos \frac{n\pi x}{L}",
        r"\lim_{n\to\infty}a_n",
        r"x\leq y\neq z",
        r"\text{if }x>0",
        r"\mathbb{R}",
        r"\begin{matrix} a & b \\ c & d \end{matrix}",
        r"\begin{aligned} x&=1 \\ y&=2 \end{aligned}",
    )

    def test_round_trip(self):
        from textconverter.renderers._omml import latex_to_omml
        for formula in self.FORMULAS:
            with self.subTest(formula=formula):
                self.assertEqual(omml_to_latex(latex_to_omml(formula)), formula)

    def test_unsupported_and_broken_formulas_raise(self):
        from textconverter.renderers._omml import UnsupportedMath, latex_to_omml
        for formula in (r"\unknowncmd{x}", r"\frac{a}{", r"a}", r"\begin{tikzpicture}x\end{tikzpicture}"):
            with self.subTest(formula=formula), self.assertRaises(UnsupportedMath):
                latex_to_omml(formula)
